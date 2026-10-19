#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
地址坐标转换工具 v2.7.1
━━━━━━━━━━━━━━━━━━━━━━
修复：批量正向/逆向编码页面空白（pages字典键名与导航不匹配）
"""

import sys, os, json, math, re, time, threading, ssl, urllib.request, urllib.parse, io, platform, subprocess
from tkinter import font as tkfont
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

# ═══════════════════════════════════════════════════════
# 延迟加载第三方库
# ═══════════════════════════════════════════════════════
_openpyxl = _PIL = None

def _get_openpyxl():
    global _openpyxl
    if _openpyxl is None:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter
        _openpyxl = (openpyxl, PatternFill, Font, Alignment, get_column_letter)
    return _openpyxl

def _get_pil():
    global _PIL
    if _PIL is None:
        from PIL import Image, ImageTk
        _PIL = (Image, ImageTk)
    return _PIL

# ═══════════════════════════════════════════════════════
# 常量配置
# ═══════════════════════════════════════════════════════
CONFIG_FILE = "addr_config.json"
ABNORMAL_COLOR = "FFFF00"

GAODE_GEO  = "https://restapi.amap.com/v3/geocode/geo"
GAODE_REGO = "https://restapi.amap.com/v3/geocode/regeo"
GAODE_STATIC = "https://restapi.amap.com/v3/staticmap"
BAIDU_GEO  = "https://api.map.baidu.com/geocoding/v3/"
BAIDU_REGO = "https://api.map.baidu.com/reverse_geocoding/v3/"
BAIDU_STATIC = "https://api.map.baidu.com/staticimage/v2"

C_PRIMARY   = "#1e293b"
C_SECONDARY = "#334155"
C_ACCENT    = "#d97706"
C_BG        = "#f1f5f9"
C_CARD      = "#ffffff"
C_TEXT      = "#0f172a"
C_TEXT_M    = "#475569"
C_TEXT_LIGHT = "#64748b"
C_BORDER    = "#e2e8f0"
C_SUCCESS   = "#059669"
C_DANGER    = "#dc2626"
C_WARNING   = "#d97706"
C_INFO      = "#2563eb"

PI = 3.14159265358979
X_PI = 52.3598775598299
EARTH_A, EARTH_EE = 6378245.0, 0.006693421622965943

COLS = {
    "serial": 1, "province": 2, "city": 3, "district": 4,
    "addr1": 5, "addr2": 6, "gcj_lng": 7, "gcj_lat": 8,
    "wgs_lng": 9, "wgs_lat": 10, "std_addr": 11,
    "rego_full": 12, "rego_div": 13, "rego_street": 14
}

GAODE_ERROR_MAP = {
    "10001": "高德 Key 无效或已过期", "10003": "高德日调用额度已用完",
    "10004": "高德 Key 类型不正确（需Web服务Key）", "10005": "高德 Key IP白名单限制",
    "10007": "高德数字签名验证失败", "20003": "高德静态地图权限不足或服务端异常",
}
BAIDU_ERROR_MAP = {
    "1": "百度服务器内部错误", "3": "百度权限校验失败", "4": "百度配额校验失败",
    "5": "百度 AK 不存在或非法", "302": "百度天配额超限", "401": "百度天配额超限",
}

LEVEL_SCORE = {
    "门址": 100, "门牌号": 100, "门址点": 100, "门牌": 100,
    "POI": 90, "兴趣点": 90, "楼宇": 90, "大厦": 90, "建筑物": 90, "购物": 90,
    "村庄": 80, "自然村": 80, "社区": 80, "村委会": 80, "住宅小区": 80,
    "道路": 70, "街道": 70, "路段": 70, "道路交叉口": 70,
    "乡镇": 60, "乡镇级": 60, "街道级": 60, "镇": 60,
    "区县": 50, "区县级": 50, "区县级别": 50,
    "城市": 30, "市级": 30, "地市级": 30, "市": 30, "省": 10, "省级": 10
}


def get_level_score(level): return LEVEL_SCORE.get(level, 20)

def get_gaode_error_msg(code, default=""):
    return GAODE_ERROR_MAP.get(str(code), default or f"高德未知错误码: {code}")

def get_baidu_error_msg(code, default=""):
    return BAIDU_ERROR_MAP.get(str(code), default or f"百度未知错误码: {code}")


# ═══════════════════════════════════════════════════════
# 工具函数
# ═══════════════════════════════════════════════════════
def clean_address(addr):
    if not addr: return ""
    addr = str(addr)
    repl = {"　": " ", "，": ",", "。": ".", "；": ";", "：": ":",
            "（": "(", "）": ")", "\r\n": "", "\n": "", "\t": ""}
    for k, v in repl.items(): addr = addr.replace(k, v)
    while "  " in addr: addr = addr.replace("  ", " ")
    return addr.strip()


def remove_parentheses(s):
    s = re.sub(r"[(（][^)）]*[)）]", "", s)
    while "  " in s: s = s.replace("  ", " ")
    return s.strip()


def get_address_levels(addr):
    levels = [addr]
    a2 = remove_parentheses(addr)
    if a2 != addr and len(a2) > 4: levels.append(a2)
    for sep in [",", " ", "-"]:
        pos = a2.rfind(sep)
        if pos > 3: levels.append(a2[:pos]); break
    m = re.search(r"(.+?(省|市))", addr)
    if m and m.group(1) not in levels: levels.append(m.group(1))
    return levels


def build_full_address(prov, city, dist, addr):
    result = ""
    for part in [prov, city, dist]:
        if part and part not in addr and part not in result: result += part
    return result + addr


def has_detail(addr):
    return bool(addr) and any(k in addr for k in ["路", "街", "号", "村", "小区", "大厦", "道", "巷", "弄", "镇", "乡", "社区"])


def check_drift(std_addr, prov, city, dist):
    if not std_addr: return True
    return any(p and p not in std_addr for p in [dist, city, prov])


def extract_core(raw_addr, prov, city, dist):
    s = raw_addr
    for p in [prov, city, dist]:
        if p: s = s.replace(p, "")
    return s.strip()


def dedup_addr(addr):
    if not addr: return addr
    for _ in range(3): addr = re.sub(r'([^\s]+?)\1+', r'\1', addr)
    addr = re.sub(r'([^\s]{2,6})\s*\1', r'\1', addr)
    return addr.strip()


def url_enc(s): return urllib.parse.quote(str(s), safe="")


def http_get(url, max_retry=3, timeout=15):
    ctx = ssl.create_default_context()
    ctx.check_hostname, ctx.verify_mode = False, ssl.CERT_NONE
    for retry in range(1, max_retry + 1):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
                data = resp.read().decode("utf-8")
                if data: return True, data
        except Exception as e:
            if retry == max_retry: return False, str(e)
            time.sleep(0.3 * retry)
    return False, ""


def json_val(js, key):
    m = re.search(r'"' + re.escape(key) + r'"\s*:\s*("?)([^"\,}\s]+)\1', js)
    return m.group(2) if m else ""


def geo_field(js, field):
    m = re.search(r'"geocodes"\s*:\s*\[\s*\{.*?"' + re.escape(field) + r'"\s*:\s*"([^"]*)"', js, re.DOTALL)
    return m.group(1) if m else ""


def regeo_field(js, field):
    m = re.search(r'"regeocode"\s*:\s*\{.*?"' + re.escape(field) + r'"\s*:\s*"([^"]*)"', js, re.DOTALL)
    return m.group(1) if m else ""


def regeo_comp(js, field):
    m = re.search(r'"addressComponent"\s*:\s*\{.*?"' + re.escape(field) + r'"\s*:\s*"([^"]*)"', js, re.DOTALL)
    return m.group(1) if m else ""


def baidu_formatted(js):
    try:
        d = json.loads(js)
        if d.get("status") in (0, "0") and isinstance(d.get("result"), dict):
            r = d["result"]
            for k in ["formatted_address", "sematic_description"]:
                v = r.get(k)
                if v and str(v).strip(): return str(v).strip()
            ac = r.get("addressComponent", {})
            parts = [str(ac.get(k, "")).strip() for k in ["province", "city", "district", "town", "street", "street_number"] if ac.get(k)]
            if parts: return "".join(parts)
    except: pass
    return ""


# ═══════════════════════════════════════════════════════
# 坐标转换
# ═══════════════════════════════════════════════════════
def _delta(lng, lat):
    dlat = -100.0 + 2.0 * lng + 3.0 * lat + 0.2 * lat * lat + 0.1 * lng * lat + 0.2 * math.sqrt(abs(lng))
    dlat += (20.0 * math.sin(6.0 * lng * PI) + 20.0 * math.sin(2.0 * lng * PI)) * 2.0 / 3.0
    dlat += (20.0 * math.sin(lat * PI) + 40.0 * math.sin(lat / 3.0 * PI)) * 2.0 / 3.0
    dlat += (160.0 * math.sin(lat / 12.0 * PI) + 320.0 * math.sin(lat * PI / 30.0)) * 2.0 / 3.0
    dlng = 300.0 + lng + 2.0 * lat + 0.1 * lng * lng + 0.1 * lng * lat + 0.1 * math.sqrt(abs(lng))
    dlng += (20.0 * math.sin(6.0 * lng * PI) + 20.0 * math.sin(2.0 * lng * PI)) * 2.0 / 3.0
    dlng += (20.0 * math.sin(lng * PI) + 40.0 * math.sin(lng / 3.0 * PI)) * 2.0 / 3.0
    dlng += (150.0 * math.sin(lng / 12.0 * PI) + 300.0 * math.sin(lng * PI / 30.0)) * 2.0 / 3.0
    return dlng, dlat


def _iter_transform(src_lng, src_lat, direction):
    lng, lat = src_lng, src_lat
    out_lng, out_lat = src_lng, src_lat
    for _ in range(3):
        dlng, dlat = _delta(lng - 105, lat - 35)
        radlat = lat / 180.0 * PI
        magic = 1 - EARTH_EE * math.sin(radlat) ** 2
        sqrtmagic = math.sqrt(magic)
        dlat = (dlat * 180.0) / ((EARTH_A * (1 - EARTH_EE)) / (magic * sqrtmagic) * PI)
        dlng = (dlng * 180.0) / (EARTH_A / sqrtmagic * math.cos(radlat) * PI)
        out_lat = src_lat + dlat * direction
        out_lng = src_lng + dlng * direction
        lat, lng = out_lat, out_lng
    return out_lng, out_lat


def wgs84_to_gcj02(lng, lat): return _iter_transform(lng, lat, 1)
def gcj02_to_wgs84(lng, lat): return _iter_transform(lng, lat, -1)


def bd09_to_gcj02(bd_lng, bd_lat):
    x, y = bd_lng - 0.0065, bd_lat - 0.006
    z = math.sqrt(x * x + y * y) - 0.00002 * math.sin(y * X_PI)
    theta = math.atan2(y, x) - 0.000003 * math.cos(x * X_PI)
    return z * math.cos(theta), z * math.sin(theta)


def gcj02_to_bd09(lng, lat):
    z = math.sqrt(lng * lng + lat * lat) + 0.00002 * math.sin(lat * X_PI)
    theta = math.atan2(lat, lng) + 0.000003 * math.cos(lng * X_PI)
    return z * math.cos(theta) + 0.0065, z * math.sin(theta) + 0.006


def convert_coords(lng, lat, src_sys, dst_sys):
    if src_sys == dst_sys: return lng, lat
    if src_sys == "wgs84": gcj_lng, gcj_lat = wgs84_to_gcj02(lng, lat)
    elif src_sys == "bd09": gcj_lng, gcj_lat = bd09_to_gcj02(lng, lat)
    else: gcj_lng, gcj_lat = lng, lat
    if dst_sys == "wgs84": return gcj02_to_wgs84(gcj_lng, gcj_lat)
    elif dst_sys == "bd09": return gcj02_to_bd09(gcj_lng, gcj_lat)
    else: return gcj_lng, gcj_lat


# ═══════════════════════════════════════════════════════
# 配置管理
# ═══════════════════════════════════════════════════════
def load_config():
    default = {
        "gaode_key": "64e7b46f67c3ceadf88394813010c230",
        "baidu_ak": "lD2wo5K150bUXVbAvcRucSRBR4YufIvj",
        "enable_baidu_fallback": True, "baidu_trigger_level": "乡镇",
        "max_retry": 3, "base_interval": 0.2, "skip_existing": True,
        "check_drift": True, "enable_cache": True, "baidu_debug": True,
        "auto_open_excel": False
    }
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                default.update(json.load(f))
        except: pass
    return default


def save_config(cfg):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


# ═══════════════════════════════════════════════════════
# 核心地理编码
# ═══════════════════════════════════════════════════════
class GeoCoder:
    def __init__(self, config):
        self.cfg = config
        self.gaode_key = config.get("gaode_key", "")
        self.baidu_ak = config.get("baidu_ak", "")
        self.enable_baidu = config.get("enable_baidu_fallback", True)
        self.trigger_level = config.get("baidu_trigger_level", "乡镇")
        self.max_retry = config.get("max_retry", 3)
        self.base_interval = config.get("base_interval", 0.2)
        self.gaode_count = self.baidu_count = 0
        self.cache, self.regeo_cache, self.coord_cache = {}, {}, {}
        self.stopped = False
        self.key_checked = False
        self.gaode_key_valid = self.baidu_key_valid = False

    def stop(self): self.stopped = True

    def init_cache(self):
        self.cache, self.regeo_cache, self.coord_cache = {}, {}, {}
        self.gaode_count = self.baidu_count = 0

    def check_keys(self):
        if self.key_checked: return True, ""
        if not self.gaode_key or len(self.gaode_key) < 10:
            return False, "高德密钥为空或格式异常"
        self.gaode_key_valid = self._check_gaode()
        if not self.gaode_key_valid:
            return True, "高德密钥验证失败，将继续尝试"
        if self.enable_baidu and self.baidu_ak:
            self.baidu_key_valid = self._check_baidu()
        self.key_checked = True
        return True, ""

    def _check_gaode(self):
        try:
            url = f"{GAODE_GEO}?address={url_enc('北京市天安门')}&output=json&key={self.gaode_key}"
            ok, resp = http_get(url, 2)
            return ok and json_val(resp, "status") == "1"
        except: return False

    def _check_baidu(self):
        try:
            url = f"{BAIDU_GEO}?address={url_enc('北京市天安门')}&output=json&ak={self.baidu_ak}"
            ok, resp = http_get(url, 2)
            return ok and json_val(resp, "status") == "0"
        except: return False

    def gaode_geocode(self, address, city=""):
        url = f"{GAODE_GEO}?address={url_enc(address)}"
        if city: url += f"&city={url_enc(city)}"
        url += f"&output=json&key={self.gaode_key}"
        ok, resp = http_get(url, self.max_retry)
        if not ok: return None, resp
        self.gaode_count += 1
        status = json_val(resp, "status")
        if status != "1":
            code = json_val(resp, "infocode")
            return None, f"高德错误[{code}]: {get_gaode_error_msg(code, json_val(resp, 'info'))}"
        return {
            "location": geo_field(resp, "location"),
            "std_addr": geo_field(resp, "formatted_address"),
            "level": geo_field(resp, "level"),
            "source": "高德"
        }, ""

    def gaode_regeocode(self, lng, lat):
        loc = f"{lng},{lat}"
        if self.cfg.get("enable_cache", True) and loc in self.regeo_cache:
            return self.regeo_cache[loc], ""
        url = f"{GAODE_REGO}?location={loc}&output=json&extensions=all&key={self.gaode_key}"
        ok, resp = http_get(url, self.max_retry)
        if not ok: return None, resp
        self.gaode_count += 1
        status = json_val(resp, "status")
        if status != "1":
            code = json_val(resp, "infocode")
            return None, f"高德逆地理错误[{code}]: {get_gaode_error_msg(code, json_val(resp, 'info'))}"
        result = {
            "full": regeo_field(resp, "formatted_address"),
            "division": f"{regeo_comp(resp, 'province')}{regeo_comp(resp, 'city')}{regeo_comp(resp, 'district')}",
            "street": f"{regeo_comp(resp, 'street')}{regeo_comp(resp, 'number')}",
            "province": regeo_comp(resp, "province"), "city": regeo_comp(resp, "city"),
            "district": regeo_comp(resp, "district"), "township": regeo_comp(resp, "township")
        }
        if self.cfg.get("enable_cache", True): self.regeo_cache[loc] = result
        return result, ""

    def baidu_geocode(self, address, city=""):
        clean = clean_address(address)
        for ch in ["&", "#", "=", "?", "%"]: clean = clean.replace(ch, "")
        clean = clean.strip()
        if not clean: return None, "地址为空"
        url = f"{BAIDU_GEO}?address={url_enc(clean)}&output=json&ak={self.baidu_ak}"
        if city: url += f"&city={url_enc(city)}"
        ok, resp = http_get(url, 2)
        if not ok: return None, resp
        self.baidu_count += 1
        try: d = json.loads(resp)
        except: return None, "JSON解析失败"
        status = d.get("status")
        if status not in (0, "0") and city and status in (1, "1"):
            url2 = f"{BAIDU_GEO}?address={url_enc(clean)}&output=json&ak={self.baidu_ak}"
            ok2, resp2 = http_get(url2, 2)
            if ok2:
                self.baidu_count += 1
                try:
                    d2 = json.loads(resp2)
                    if d2.get("status") in (0, "0"): d, status = d2, d2.get("status")
                    else: return None, f"百度错误: {get_baidu_error_msg(str(d.get('status','')), '未知错误')}"
                except: return None, "二次JSON解析失败"
            else: return None, f"百度错误: {get_baidu_error_msg(str(d.get('status','')), '未知错误')}"
        if d.get("status") not in (0, "0"):
            return None, f"百度错误: {get_baidu_error_msg(str(d.get('status','')), '未知错误')}"
        result = d.get("result", {})
        if not isinstance(result, dict): return None, "无result"
        loc = result.get("location", {})
        try: bd_lng, bd_lat = float(loc.get("lng", 0)), float(loc.get("lat", 0))
        except: return None, "坐标格式错误"
        if bd_lng == 0 or bd_lat == 0: return None, "坐标为0"
        gcj_lng, gcj_lat = bd09_to_gcj02(bd_lng, bd_lat)
        level = result.get("level", "未知")
        std_addr = baidu_formatted(resp)
        if not std_addr:
            ck = (round(bd_lng, 6), round(bd_lat, 6))
            if ck in self.coord_cache: std_addr = self.coord_cache[ck]
            else:
                try:
                    re_url = f"{BAIDU_REGO}?ak={self.baidu_ak}&output=json&coordtype=bd09ll&location={bd_lat},{bd_lng}"
                    ok_re, resp_re = http_get(re_url, 2)
                    if ok_re:
                        self.baidu_count += 1
                        d_re = json.loads(resp_re)
                        if d_re.get("status") in (0, "0"):
                            std_addr = baidu_formatted(resp_re)
                            if std_addr: self.coord_cache[ck] = std_addr
                except: pass
        if std_addr and level:
            try:
                ac = result.get("addressComponent", {})
                ls = get_level_score(level)
                if ls <= get_level_score("区县"):
                    parts = [str(ac.get(k, "")).strip() for k in ["province", "city", "district"] if ac.get(k)]
                    if parts: std_addr = "".join(parts)
                elif ls <= get_level_score("乡镇"):
                    parts = [str(ac.get(k, "")).strip() for k in ["province", "city", "district", "town"] if ac.get(k)]
                    if parts: std_addr = "".join(parts)
            except: pass
        std_addr = dedup_addr(std_addr) or clean
        return {
            "location": f"{gcj_lng},{gcj_lat}", "std_addr": std_addr,
            "level": level, "source": "百度"
        }, ""

    def resolve(self, addr_str, prov_str, city_str, dist_str):
        baidu_log = ""
        if self.stopped: return None, False, "用户取消", False, baidu_log
        if not addr_str: return None, True, "详细地址为空", False, baidu_log
        api_city = "".join(filter(None, [prov_str, city_str, dist_str]))
        gaode_city = city_str if "市" in city_str else ""
        baidu_city = gaode_city
        query = build_full_address(prov_str, city_str, dist_str, addr_str)
        clean_q = clean_address(query)
        cache_key = f"{api_city}|{addr_str}".lower().replace(" ", "").replace("，", "").replace("。", "")
        if self.cfg.get("enable_cache", True) and cache_key in self.cache:
            c = self.cache[cache_key]
            return c, c.get("is_abnormal", False), "", False, baidu_log
        has_detail_flag = has_detail(addr_str)
        levels = get_address_levels(clean_q)
        gcj_lng, gcj_lat, std_addr, level, source = 0.0, 0.0, "", "", "高德"
        is_drift, is_abnormal, err_msg, location = False, False, "", ""
        for lvl_addr in levels:
            if self.stopped: return None, False, "用户取消", False, baidu_log
            result, err = self.gaode_geocode(lvl_addr, gaode_city)
            if result:
                location = result["location"]
                std_addr = result["std_addr"]
                level = result["level"]
                source = result["source"]
                if location and "," in location and get_level_score(level) >= get_level_score("街道"):
                    break
            if err:
                for fatal_code in ["10001", "10003", "10004", "10005", "10007"]:
                    if fatal_code in err:
                        return None, True, f"高德全局错误：{err}", True, baidu_log
            time.sleep(self.base_interval)
        if not location:
            err_msg, is_abnormal = "高德无匹配结果", True
        else:
            try:
                parts = location.split(",")
                gcj_lng, gcj_lat = float(parts[0]), float(parts[1])
            except: gcj_lng, gcj_lat = 0.0, 0.0
        if self.cfg.get("check_drift", True) and api_city and std_addr:
            is_drift = check_drift(std_addr, prov_str, city_str, dist_str)
            if is_drift: is_abnormal = True
        if self.enable_baidu and location and self.baidu_ak and self.baidu_key_valid:
            gaode_score = get_level_score(level)
            trigger_score = get_level_score(self.trigger_level)
            need_baidu = is_drift or (gaode_score <= trigger_score and has_detail_flag)
            if need_baidu:
                baidu_log = f"[百度Fallback] 高德精度不足({level})，尝试百度解析...\n"
                bd_result, bd_err = self.baidu_geocode(query, baidu_city)
                if bd_result:
                    bd_score = get_level_score(bd_result["level"])
                    bd_drift = check_drift(bd_result["std_addr"], prov_str, city_str, dist_str)
                    if bd_score > gaode_score and not bd_drift:
                        parts = bd_result["location"].split(",")
                        gcj_lng, gcj_lat = float(parts[0]), float(parts[1])
                        std_addr, level, source = bd_result["std_addr"], bd_result["level"], "百度"
                        is_drift, is_abnormal = False, False
                        baidu_log += f"[百度Fallback] 百度提升成功: {std_addr} [精度:{level}]\n"
                    else:
                        baidu_log += f"[百度Fallback] 百度未提升({bd_result.get('level','')})，保留高德结果\n"
                        if self.cfg.get("baidu_debug", True): std_addr += f" [百度未提升:{bd_result.get('level','')}]"
                else:
                    baidu_log += f"[百度Fallback] 百度解析失败: {bd_err}\n"
                    if self.cfg.get("baidu_debug", True): std_addr += f" [百度失败:{bd_err}]"
                    is_abnormal = True
        if gcj_lng and gcj_lat:
            if self.cfg.get("enable_cache", True) and not is_abnormal and get_level_score(level) > get_level_score("区县"):
                self.cache[cache_key] = {
                    "gcj_lng": gcj_lng, "gcj_lat": gcj_lat, "std_addr": std_addr,
                    "level": level, "source": source, "is_drift": is_drift, "is_abnormal": is_abnormal
                }
            return {
                "gcj_lng": gcj_lng, "gcj_lat": gcj_lat, "std_addr": std_addr,
                "level": level, "source": source, "is_drift": is_drift, "is_abnormal": is_abnormal
            }, is_abnormal, err_msg or ("解析异常" if is_abnormal else ""), False, baidu_log
        return None, True, err_msg or "双平台均无法解析", False, baidu_log


# ═══════════════════════════════════════════════════════
# Excel 处理
# ═══════════════════════════════════════════════════════
class ExcelProcessor:
    def __init__(self, filepath):
        openpyxl, PatternFill, Font, Alignment, get_column_letter = _get_openpyxl()
        self.openpyxl, self.PatternFill, self.Font, self.Alignment = openpyxl, PatternFill, Font, Alignment
        self.get_column_letter = get_column_letter
        self.filepath = filepath
        self.wb = openpyxl.load_workbook(filepath)
        if "Sheet1" not in self.wb.sheetnames: raise ValueError("Excel 中未找到 Sheet1")
        self.ws = self.wb["Sheet1"]
        self.yellow_fill = PatternFill(start_color=ABNORMAL_COLOR, end_color=ABNORMAL_COLOR, fill_type="solid")
        self.none_fill = PatternFill(fill_type=None)

    def get_last_row(self, col=COLS["addr1"]):
        for row in range(self.ws.max_row, 1, -1):
            val = self.ws.cell(row=row, column=col).value
            if val is not None and str(val).strip(): return row
        return 1

    def is_normal(self, row):
        lng = self.ws.cell(row=row, column=COLS["gcj_lng"]).value
        addr = str(self.ws.cell(row=row, column=COLS["std_addr"]).value or "")
        if not lng or not isinstance(lng, (int, float)) or lng == 0: return False
        return not any(k in addr for k in ["失败", "错误", "漂移", "无匹配", "空地址", "无效", "超出范围"])

    def is_abnormal(self, row):
        if self.is_normal(row): return False
        addr = str(self.ws.cell(row=row, column=COLS["std_addr"]).value or "")
        return bool(addr) and any(k in addr for k in ["失败", "错误", "无匹配", "空地址", "无效", "超出范围"])

    def get_abnormal_rows(self, start_row=2):
        abnormal = []
        last = self.get_last_row()
        headers = [self.ws.cell(row=1, column=c).value for c in range(1, self.ws.max_column + 1)]
        for r in range(start_row, last + 1):
            if self.is_abnormal(r):
                row_data = {headers[c - 1] if headers[c - 1] else f"列{c}": self.ws.cell(row=r, column=c).value
                            for c in range(1, self.ws.max_column + 1)}
                row_data["_原始行号"] = r
                abnormal.append(row_data)
        return abnormal, headers

    def export_abnormal_rows(self, output_path, start_row=2):
        abnormal, headers = self.get_abnormal_rows(start_row)
        if not abnormal: return False, "没有异常行需要导出"
        try:
            wb = self.openpyxl.Workbook()
            ws = wb.active; ws.title = "异常行"
            out_headers = ["原始行号"] + [h for h in headers if h]
            for col, h in enumerate(out_headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = self.Font(bold=True, color="FFFFFF")
                cell.fill = self.PatternFill(start_color="dc3545", end_color="dc3545", fill_type="solid")
                cell.alignment = self.Alignment(horizontal="center", vertical="center")
            for r_idx, row_data in enumerate(abnormal, 2):
                ws.cell(row=r_idx, column=1, value=row_data.get("_原始行号", ""))
                for c_idx, h in enumerate(out_headers[1:], 2):
                    ws.cell(row=r_idx, column=c_idx, value=row_data.get(h, ""))
            for i, h in enumerate(out_headers, 1):
                w = min(max(len(str(h)) * 1.5 + 2, 10), 40)
                ws.column_dimensions[self.get_column_letter(i)].width = w
            wb.save(output_path)
            return True, f"已导出 {len(abnormal)} 条异常行"
        except Exception as e:
            return False, f"导出失败: {e}"

    def clear_forward(self, row):
        for c in range(COLS["gcj_lng"], COLS["std_addr"] + 1):
            self.ws.cell(row=row, column=c).value = None
            self.ws.cell(row=row, column=c).fill = self.none_fill

    def write_forward(self, row, result):
        if not result:
            for c in range(COLS["gcj_lng"], COLS["std_addr"] + 1):
                self.ws.cell(row=row, column=c).fill = self.yellow_fill
            self.ws.cell(row=row, column=COLS["gcj_lng"]).value = "无匹配"
            self.ws.cell(row=row, column=COLS["gcj_lat"]).value = "无匹配"
            self.ws.cell(row=row, column=COLS["wgs_lng"]).value = "无匹配"
            self.ws.cell(row=row, column=COLS["wgs_lat"]).value = "无匹配"
            self.ws.cell(row=row, column=COLS["std_addr"]).value = "解析失败"
            return
        gcj_lng, gcj_lat = result["gcj_lng"], result["gcj_lat"]
        self.ws.cell(row=row, column=COLS["gcj_lng"]).value = gcj_lng
        self.ws.cell(row=row, column=COLS["gcj_lat"]).value = gcj_lat
        if 73 <= gcj_lng <= 135 and 3 <= gcj_lat <= 53:
            wgs_lng, wgs_lat = gcj02_to_wgs84(gcj_lng, gcj_lat)
            self.ws.cell(row=row, column=COLS["wgs_lng"]).value = round(wgs_lng, 6)
            self.ws.cell(row=row, column=COLS["wgs_lat"]).value = round(wgs_lat, 6)
        else:
            self.ws.cell(row=row, column=COLS["wgs_lng"]).value = "超出范围"
            self.ws.cell(row=row, column=COLS["wgs_lat"]).value = "超出范围"
            result["is_abnormal"] = True
        final = f"{result['std_addr']} [精度:{result['level']}][来源:{result['source']}]"
        if result.get("is_drift"): final += " 疑似漂移"
        self.ws.cell(row=row, column=COLS["std_addr"]).value = final
        fill = self.yellow_fill if result.get("is_abnormal") else self.none_fill
        for c in range(COLS["gcj_lng"], COLS["std_addr"] + 1):
            self.ws.cell(row=row, column=c).fill = fill

    def clear_reverse(self, row):
        for c in range(COLS["rego_full"], COLS["rego_street"] + 1):
            self.ws.cell(row=row, column=c).value = None

    def write_reverse(self, row, result):
        if not result:
            self.ws.cell(row=row, column=COLS["rego_full"]).value = "解析失败"
            return
        self.ws.cell(row=row, column=COLS["rego_full"]).value = result.get("full", "")
        self.ws.cell(row=row, column=COLS["rego_div"]).value = result.get("division", "")
        self.ws.cell(row=row, column=COLS["rego_street"]).value = result.get("street", "")

    def save(self): self.wb.save(self.filepath)


# ═══════════════════════════════════════════════════════
# 精致化 UI 组件
# ═══════════════════════════════════════════════════════
class SButton(tk.Button):
    def __init__(self, master, text, command, btn_type="primary", **kw):
        colors = {
            "primary":   (C_ACCENT, "#fff", "#b45309"),
            "secondary": (C_INFO,   "#fff", "#1d4ed8"),
            "success":   (C_SUCCESS, "#fff", "#047857"),
            "danger":    (C_DANGER, "#fff", "#b91c1c"),
            "warning":   ("#f59e0b", "#fff", "#d97706"),
        }
        bg, fg, active = colors.get(btn_type, colors["primary"])
        super().__init__(master, text=text, command=command,
                         font=("微软雅黑", 10, "bold"), bg=bg, fg=fg, bd=0,
                         padx=18, pady=8, cursor="hand2", relief=tk.FLAT,
                         activebackground=active, activeforeground="#fff", **kw)
        self._normal_bg = bg
        self.bind("<Enter>", lambda e: self.config(bg=active))
        self.bind("<Leave>", lambda e: self.config(bg=self._normal_bg))


class Card(tk.Frame):
    def __init__(self, master, title=None, accent=C_ACCENT, **kw):
        super().__init__(master, bg=C_CARD, **kw)
        self.configure(highlightbackground=C_BORDER, highlightthickness=1)
        if accent:
            tk.Frame(self, bg=accent, height=3).pack(fill=tk.X)
        if title:
            tf = tk.Frame(self, bg=C_CARD)
            tf.pack(fill=tk.X, padx=14, pady=(10, 0))
            tk.Label(tf, text=title, font=("微软雅黑", 12, "bold"), bg=C_CARD, fg=C_TEXT).pack(anchor="w")
            tk.Frame(self, bg=C_BORDER, height=1).pack(fill=tk.X, padx=14, pady=8)


class CollapsibleSection(tk.Frame):
    def __init__(self, master, title, expanded=False, title_font=None,
                 title_bg=C_CARD, title_fg=C_ACCENT, arrow_color=C_ACCENT, **kw):
        super().__init__(master, bg=C_CARD, **kw)
        self.expanded = expanded
        self.header = tk.Frame(self, bg=title_bg, cursor="hand2")
        self.header.pack(fill=tk.X)
        self.header.bind("<Button-1>", self._toggle)
        self.arrow = tk.Label(self.header, text="▼" if expanded else "▶", font=("微软雅黑", 9, "bold"),
                              bg=title_bg, fg=arrow_color, width=2)
        self.arrow.pack(side=tk.LEFT, padx=(10, 4))
        self.arrow.bind("<Button-1>", self._toggle)
        tf = title_font or ("微软雅黑", 12, "bold")
        self.title_lbl = tk.Label(self.header, text=title, font=tf, bg=title_bg, fg=title_fg)
        self.title_lbl.pack(side=tk.LEFT, padx=(0, 8), pady=8)
        self.title_lbl.bind("<Button-1>", self._toggle)
        self.sep = tk.Frame(self, bg=C_BORDER, height=1)
        self.sep.pack(fill=tk.X, padx=10)
        self.content = tk.Frame(self, bg=C_CARD)
        if expanded: self.content.pack(fill=tk.X, padx=10, pady=(6, 10))
        for w in (self.header, self.arrow, self.title_lbl):
            w.bind("<Enter>", lambda e: self.header.config(bg="#f8fafc"))
            w.bind("<Leave>", lambda e: self.header.config(bg=title_bg))

    def _toggle(self, event=None):
        self.expanded = not self.expanded
        self.arrow.config(text="▼" if self.expanded else "▶")
        if self.expanded: self.content.pack(fill=tk.X, padx=10, pady=(6, 10))
        else: self.content.pack_forget()

    def get_content_frame(self): return self.content


class PreviewManager:
    def __init__(self, parent_frame, notebook, tab_text="📋 数据预览"):
        self.notebook = notebook
        self.page_size = 20
        self.data = []
        self.headers = []
        self.page = 0
        self.total_pages = 0

        self.frame = tk.Frame(notebook, bg=C_CARD)
        notebook.add(self.frame, text=f"  {tab_text}  ")

        toolbar = tk.Frame(self.frame, bg=C_CARD)
        toolbar.pack(fill=tk.X, padx=10, pady=(8, 4))
        self.info_lbl = tk.Label(toolbar, text="请先导入 Excel 文件", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT_LIGHT)
        self.info_lbl.pack(side=tk.LEFT)

        tree_frame = tk.Frame(self.frame, bg=C_CARD)
        tree_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 4))
        self.tree = ttk.Treeview(tree_frame, show="headings", selectmode="browse", height=16)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        vsb.grid(row=0, column=1, sticky="ns")
        hsb = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        hsb.grid(row=1, column=0, sticky="ew")
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.bind("<MouseWheel>", lambda e: self.tree.yview_scroll(int(-1 * (e.delta / 120)), "units"))
        self.tree.bind("<Shift-MouseWheel>", lambda e: self.tree.xview_scroll(int(-1 * (e.delta / 120)), "units"))

        page_frame = tk.Frame(self.frame, bg=C_CARD)
        page_frame.pack(fill=tk.X, padx=10, pady=(0, 8))
        self.page_lbl = tk.Label(page_frame, text="第 0 / 0 页", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT)
        self.page_lbl.pack(side=tk.LEFT)
        tk.Button(page_frame, text="◀ 上一页", font=("微软雅黑", 9), bg=C_BORDER, fg=C_TEXT,
                  bd=0, padx=12, pady=4, cursor="hand2", relief=tk.FLAT,
                  command=self.prev_page).pack(side=tk.LEFT, padx=(12, 4))
        tk.Button(page_frame, text="下一页 ▶", font=("微软雅黑", 9), bg=C_BORDER, fg=C_TEXT,
                  bd=0, padx=12, pady=4, cursor="hand2", relief=tk.FLAT,
                  command=self.next_page).pack(side=tk.LEFT, padx=4)

    def load_data(self, headers, data):
        self.headers = headers
        self.data = data
        self.page = 0
        self.total_pages = max(1, (len(data) + self.page_size - 1) // self.page_size) if data else 0
        self.refresh_page()
        self.info_lbl.config(text=f"共 {len(data)} 行数据 | {len(headers)} 列（预览前 {min(len(data), 200)} 行）")

    def refresh_page(self):
        for col in self.tree["columns"]:
            self.tree.heading(col, text="")
        self.tree.delete(*self.tree.get_children())
        if not self.data:
            self.tree["columns"] = ("提示",)
            self.tree.heading("提示", text="暂无数据，请先导入 Excel 文件")
            self.tree.column("提示", width=400)
            self.page_lbl.config(text="第 0 / 0 页")
            return
        display_headers = ["序号"] + self.headers
        self.tree["columns"] = display_headers
        self.tree.column("序号", width=55, anchor="center")
        self.tree.heading("序号", text="序号")
        for h in display_headers[1:]:
            self.tree.column(h, width=120, anchor="w", minwidth=80, stretch=False)
            self.tree.heading(h, text=h)
        start_idx = self.page * self.page_size
        end_idx = min(start_idx + self.page_size, len(self.data))
        for i in range(start_idx, end_idx):
            row = self.data[i]
            values = [i + 1]
            for h in display_headers[1:]:
                val = row.get(h, "")
                if len(str(val)) > 30: val = str(val)[:27] + "..."
                values.append(val)
            self.tree.insert("", tk.END, values=values)
        self.page_lbl.config(text=f"第 {self.page + 1} / {self.total_pages} 页  （共 {len(self.data)} 行）")
        self.tree.update_idletasks()

    def next_page(self):
        if self.page < self.total_pages - 1:
            self.page += 1; self.refresh_page()

    def prev_page(self):
        if self.page > 0:
            self.page -= 1; self.refresh_page()


# ═══════════════════════════════════════════════════════
# 批量处理核心
# ═══════════════════════════════════════════════════════
class BatchProcessor:
    def __init__(self, app):
        self.app = app
        self.geo = None
        self.stop_flag = False

    def stop(self):
        self.stop_flag = True
        if self.geo: self.geo.stop()

    def init(self, config):
        self.geo = GeoCoder(config)
        self.stop_flag = False
        return self.geo.check_keys()

    def process_forward(self, proc, start_row, last_row, addr_idx, addr2_idx, city_idx,
                        skip_existing, process_abnormal, log_callback, prog_callback):
        total = last_row - start_row + 1
        success = fail = skip = addr2_used = drift_cnt = baidu_up = 0
        for i in range(start_row, last_row + 1):
            if self.stop_flag:
                log_callback("\n⚠ 用户取消处理\n", "warning"); break
            pv = i - start_row + 1
            prog_callback(pv, total)
            if skip_existing and proc.is_normal(i):
                skip += 1; log_callback(f"第 {i} 行：跳过（已有结果）\n", "skip"); continue
            if process_abnormal and not proc.is_abnormal(i):
                skip += 1; continue
            proc.clear_forward(i)
            city = str(proc.ws.cell(row=i, column=city_idx).value or "").strip()
            addr1 = str(proc.ws.cell(row=i, column=addr_idx).value or "").strip()
            addr2 = str(proc.ws.cell(row=i, column=addr2_idx).value or "").strip()
            if not addr1:
                proc.write_forward(i, None); fail += 1
                log_callback(f"第 {i} 行：地址为空\n", "error"); continue
            result, is_abnormal_flag, err_msg, fatal, baidu_log = self.geo.resolve(addr1, "", city, "")
            if fatal:
                log_callback(f"致命错误：{err_msg}，已终止\n", "error"); break
            if result and not is_abnormal_flag and result.get("source") == "高德" and addr2:
                core = extract_core(addr1, "", city, "")
                if core and core.lower() not in result["std_addr"].lower():
                    is_abnormal_flag = True; result["is_abnormal"] = True; result["std_addr"] += " [内容不匹配]"
            if (not result) or is_abnormal_flag:
                if addr2:
                    result2, ab2, err2, fatal2, _ = self.geo.resolve(addr2, "", city, "")
                    if fatal2:
                        log_callback(f"致命错误：{err2}，已终止\n", "error"); break
                    if result2 and not ab2:
                        result = result2; addr2_used += 1
                        if result.get("is_drift"): drift_cnt += 1
                        if result.get("source") == "百度": baidu_up += 1
                    else: result = None
                else: result = None
            proc.write_forward(i, result)
            if result: success += 1; status, tag = "成功", "success"
            else: fail += 1; status, tag = "失败", "error"
            if baidu_log: log_callback(baidu_log, "baidu")
            log_callback(f"第 {i} 行：{status}\n", tag)
            time.sleep(self.geo.base_interval)
        return {
            "total": total, "success": success, "fail": fail, "skip": skip,
            "addr2_used": addr2_used, "drift_cnt": drift_cnt, "baidu_up": baidu_up,
            "gaode_count": self.geo.gaode_count, "baidu_count": self.geo.baidu_count
        }

    def process_reverse(self, proc, start_row, last_row, gcj_lng_idx, gcj_lat_idx,
                        wgs_lng_idx, wgs_lat_idx, skip_existing, log_callback, prog_callback):
        total = last_row - start_row + 1
        success = fail = skip = 0
        global_error = False
        for i in range(start_row, last_row + 1):
            if self.stop_flag:
                log_callback("\n⚠ 用户取消处理\n", "warning"); break
            pv = i - start_row + 1
            prog_callback(pv, total)
            if skip_existing:
                existing = str(proc.ws.cell(row=i, column=COLS["rego_full"]).value or "").strip()
                if existing and "失败" not in existing and "错误" not in existing:
                    skip += 1; log_callback(f"第 {i} 行：跳过\n"); continue
            proc.clear_reverse(i)
            gcj_lng = proc.ws.cell(row=i, column=gcj_lng_idx).value
            gcj_lat = proc.ws.cell(row=i, column=gcj_lat_idx).value
            try: gcj_lng = float(gcj_lng) if gcj_lng else 0; gcj_lat = float(gcj_lat) if gcj_lat else 0
            except: gcj_lng, gcj_lat = 0, 0
            if not (gcj_lng and gcj_lat and 73 <= gcj_lng <= 135 and 3 <= gcj_lat <= 53):
                wgs_lng = proc.ws.cell(row=i, column=wgs_lng_idx).value if wgs_lng_idx > 0 else 0
                wgs_lat = proc.ws.cell(row=i, column=wgs_lat_idx).value if wgs_lat_idx > 0 else 0
                try: wgs_lng = float(wgs_lng) if wgs_lng else 0; wgs_lat = float(wgs_lat) if wgs_lat else 0
                except: wgs_lng, wgs_lat = 0, 0
                if wgs_lng and wgs_lat and 73 <= wgs_lng <= 135 and 3 <= wgs_lat <= 53:
                    gcj_lng, gcj_lat = wgs84_to_gcj02(wgs_lng, wgs_lat)
                else:
                    proc.write_reverse(i, None); fail += 1
                    log_callback(f"第 {i} 行：无有效坐标\n"); continue
            if not (gcj_lng and gcj_lat):
                proc.ws.cell(row=i, column=COLS["rego_full"]).value = "无有效坐标"; fail += 1
                log_callback(f"第 {i} 行：无有效坐标\n"); continue
            result, err = self.geo.gaode_regeocode(gcj_lng, gcj_lat)
            if not result:
                proc.ws.cell(row=i, column=COLS["rego_full"]).value = f"失败：{err}"
                if "10001" in err or "10003" in err:
                    global_error = True; fail += 1
                    log_callback(f"第 {i} 行：全局错误 {err}，已终止\n"); break
                fail += 1; log_callback(f"第 {i} 行：失败 - {err}\n")
            else:
                proc.write_reverse(i, result); success += 1
                log_callback(f"第 {i} 行：成功\n")
            time.sleep(self.geo.base_interval)
        return {
            "total": total, "success": success, "fail": fail, "skip": skip,
            "global_error": global_error, "gaode_count": self.geo.gaode_count
        }


# ═══════════════════════════════════════════════════════
# 主应用
# ═══════════════════════════════════════════════════════
class App:
    def __init__(self, root):
        self.root = root
        self.root.title("地址坐标转换工具 v2.7.1")
        self.root.geometry("1400x900")
        self.root.minsize(1320, 820)
        self.root.state("zoomed")
        self.config = load_config()
        self.geo = GeoCoder(self.config)
        self.running = False
        self.stop_flag = False
        self.default_font = tkfont.Font(family="微软雅黑", size=10)
        self.root.option_add("*Font", self.default_font)
        self._last_result = {}
        self.batch_processor = BatchProcessor(self)
        self._build_ui()
        self.show_page("home")

    def _build_ui(self):
        main = tk.Frame(self.root, bg=C_BG)
        main.pack(fill=tk.BOTH, expand=True)

        nav_w = 210
        nav = tk.Frame(main, bg=C_PRIMARY, width=nav_w)
        nav.pack(side=tk.LEFT, fill=tk.Y)
        nav.pack_propagate(False)

        lf = tk.Frame(nav, bg=C_PRIMARY, height=120)
        lf.pack(fill=tk.X, pady=(35, 25))
        cvs = tk.Canvas(lf, width=52, height=52, bg=C_PRIMARY, highlightthickness=0)
        cvs.pack()
        cvs.create_oval(4, 4, 48, 48, fill=C_ACCENT, outline="")
        cvs.create_text(26, 26, text="地", font=("微软雅黑", 20, "bold"), fill=C_PRIMARY)
        tk.Label(lf, text="地址转换", font=("微软雅黑", 22, "bold"), bg=C_PRIMARY, fg=C_ACCENT).pack(pady=(12, 2))
        tk.Label(lf, text="标准化工具", font=("微软雅黑", 11), bg=C_PRIMARY, fg="#8892b0").pack()

        self.nav_btns = {}
        nf = tk.Frame(nav, bg=C_PRIMARY)
        nf.pack(fill=tk.X, padx=12, pady=10)
        nav_items = [
            ("home", "首页", "🏠"), ("single", "单条查询", "📍"),
            ("batch_fwd", "批量正向编码", "📊"), ("batch_rev", "批量逆向编码", "🔄"),
            ("converter", "坐标转换", "🧭"), ("settings", "API 设置", "⚙️")
        ]
        for key, text, icon in nav_items:
            bf = tk.Frame(nf, bg=C_PRIMARY)
            bf.pack(fill=tk.X, pady=2)
            btn = tk.Button(bf, text=f"{icon}   {text}", font=("微软雅黑", 11),
                           bg=C_PRIMARY, fg="#94a3b8", activebackground=C_SECONDARY,
                           activeforeground=C_ACCENT, bd=0, padx=14, pady=11,
                           anchor="w", cursor="hand2", relief=tk.FLAT,
                           command=lambda k=key: self.show_page(k))
            btn.pack(fill=tk.X)
            btn.bind("<Enter>", lambda e, b=btn: b.config(bg=C_SECONDARY, fg="#e2e8f0"))
            btn.bind("<Leave>", lambda e, b=btn: b.config(bg=C_PRIMARY, fg="#94a3b8"))
            self.nav_btns[key] = btn

        tk.Frame(nav, bg=C_PRIMARY).pack(expand=True)
        tk.Label(nav, text="v2.7.1", font=("微软雅黑", 10), bg=C_PRIMARY, fg="#64748b").pack(pady=14)

        self.content = tk.Frame(main, bg=C_BG)
        self.content.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=18, pady=18)

        self.status_var = tk.StringVar(value="就绪")
        sb = tk.Frame(main, bg=C_PRIMARY, height=32)
        sb.pack(side=tk.BOTTOM, fill=tk.X)
        sb.pack_propagate(False)
        tk.Label(sb, textvariable=self.status_var, font=("微软雅黑", 10), bg=C_PRIMARY, fg="#94a3b8").pack(side=tk.LEFT, padx=18, pady=4)

        self.pages = {}
        self._build_home()
        self._build_single()
        self._build_batch("fwd")
        self._build_batch("rev")
        self._build_converter()
        self._build_settings()

    def _nav_style(self, active):
        for k, btn in self.nav_btns.items():
            if k == active:
                btn.config(bg=C_SECONDARY, fg=C_ACCENT, font=("微软雅黑", 11, "bold"))
            else:
                btn.config(bg=C_PRIMARY, fg="#94a3b8", font=("微软雅黑", 11))

    def show_page(self, key):
        self._nav_style(key)
        for p in self.pages.values(): p.pack_forget()
        self.pages[key].pack(fill=tk.BOTH, expand=True)
        self.current = key

    def _make_page(self, parent):
        return tk.Frame(parent, bg=C_BG)

    def _get_col_idx(self, headers, col_name, default_col):
        if not col_name: return default_col
        try: return headers.index(col_name) + 1
        except ValueError: return default_col

    def _config_log_tags(self, log_widget):
        for tag, color in [("success", "#059669"), ("error", "#dc2626"), ("warning", "#d97706"),
                           ("info", "#2563eb"), ("skip", "#64748b"), ("baidu", "#ca8a04")]:
            log_widget.tag_config(tag, foreground=color)

    def _log_batch(self, direction, msg, tag="normal"):
        log = getattr(self, f"{direction}_log")
        log.insert(tk.END, msg, tag)
        log.see(tk.END)

    # ═══════════════════════════════════════════════════════
    # 首页
    # ═══════════════════════════════════════════════════════
    def _build_home(self):
        frame = self._make_page(self.content)
        self.pages["home"] = frame

        title_frame = tk.Frame(frame, bg=C_BG)
        title_frame.pack(fill=tk.X, pady=(0, 4))
        tk.Label(title_frame, text="欢迎使用地址转换工具", font=("微软雅黑", 24, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor="w")
        tk.Label(title_frame, text="基于高德/百度地图 API 的地址智能解析与坐标转换",
                 font=("微软雅黑", 12), bg=C_BG, fg=C_TEXT_LIGHT).pack(anchor="w", pady=(4, 0))

        cf = tk.Frame(frame, bg=C_BG)
        cf.pack(fill=tk.X, pady=(12, 8))
        card_items = [
            ("📍 单条查询", "快速查询单个地址的 GCJ-02 / WGS-84 坐标\n支持地址查询与经纬度逆编码", C_INFO),
            ("📊 批量正向编码", "从 Excel 读取地址，批量获取\n标准化结果与坐标", C_SUCCESS),
            ("🔄 批量逆向编码", "从 Excel 读取坐标，批量解析\n为标准地址信息", C_WARNING),
            ("🧭 坐标转换", "GCJ-02 / WGS-84 / BD-09\n三系坐标互转，支持单点与批量", "#7c3aed"),
            ("⚙️ 双平台配置", "支持高德 + 百度 API\n智能 Fallback 自动切换", C_DANGER)
        ]
        for i, (title, desc, color) in enumerate(card_items):
            c = tk.Frame(cf, bg=C_CARD, highlightbackground=C_BORDER, highlightthickness=1)
            c.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0 if i == 0 else 8, 0))
            tk.Frame(c, bg=color, height=4).pack(fill=tk.X)
            tk.Label(c, text=title, font=("微软雅黑", 12, "bold"), bg=C_CARD, fg=C_TEXT).pack(anchor="w", padx=12, pady=(10, 4))
            tk.Label(c, text=desc, font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT_LIGHT, justify=tk.LEFT).pack(anchor="w", padx=12, pady=(0, 10))

        steps = CollapsibleSection(frame, title="▎使用步骤", expanded=True,
                                    title_font=("微软雅黑", 13, "bold"), title_fg=C_ACCENT)
        steps.pack(fill=tk.X, pady=(10, 8))
        sc = steps.get_content_frame()
        for step in [
            "1. 点击左侧「API 设置」，配置高德/百度 API Key（自动保存到 addr_config.json）",
            "2. 准备 Excel 文件，确保包含 Sheet1，按下方表格格式填写数据",
            "3. 选择「批量正向编码」或「批量逆向编码」，选择文件后开始处理",
            "4. 结果自动写入 Excel，异常行黄色高亮；支持断点续传与失败重试"
        ]:
            tk.Label(sc, text=step, font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT_M).pack(anchor="w", padx=6, pady=3)

        fmt = CollapsibleSection(frame, title="▎Excel 数据格式说明", expanded=True,
                                  title_font=("微软雅黑", 13, "bold"), title_fg=C_SUCCESS)
        fmt.pack(fill=tk.X, pady=(0, 8))
        fc = fmt.get_content_frame()
        dl_btn = tk.Button(fmt.header, text="下载模板", font=("微软雅黑", 10, "bold"), bg=C_ACCENT, fg="#fff",
                           bd=0, padx=14, pady=4, cursor="hand2", relief=tk.FLAT, command=self.download_template)
        dl_btn.pack(side=tk.RIGHT, padx=14)
        dl_btn.bind("<Enter>", lambda e: dl_btn.config(bg="#b45309"))
        dl_btn.bind("<Leave>", lambda e: dl_btn.config(bg=C_ACCENT))

        tk.Label(fc, text="▎正向编码（地址 → 坐标）", font=("微软雅黑", 11, "bold"), bg=C_CARD, fg=C_SUCCESS).pack(anchor="w", padx=6, pady=(6, 4))
        self._table(fc, [
            ["序号", "省(选填)", "市(选填)", "县/区(选填)", "地址1(必填)", "地址2(选填)",
             "GCJ-02\n经度", "GCJ-02\n纬度", "WGS-84\n经度", "WGS-84\n纬度", "标准地址"],
            ["输入", "输入", "输入", "输入", "输入", "输入", "输出", "输出", "输出", "输出", "输出"]
        ], [6, 8, 8, 8, 16, 14, 10, 10, 10, 10, 24])

        tk.Label(fc, text="▎逆向编码（坐标 → 地址）", font=("微软雅黑", 11, "bold"), bg=C_CARD, fg=C_WARNING).pack(anchor="w", padx=6, pady=(14, 4))
        self._table(fc, [
            ["GCJ-02\n经度", "GCJ-02\n纬度", "WGS-84\n经度", "WGS-84\n纬度",
             "逆地理\n完整地址", "省 - 市 - 区", "街道\n门牌号"],
            ["输入", "输入", "输入(备选)", "输入(备选)", "输出", "输出", "输出"]
        ], [11, 11, 11, 11, 18, 18, 12])

    def _table(self, parent, rows, widths):
        tf = tk.Frame(parent, bg=C_CARD)
        tf.pack(fill=tk.X, padx=4, pady=4)
        for r_idx, row in enumerate(rows):
            if r_idx > 0: tk.Frame(tf, bg=C_BORDER, height=1).pack(fill=tk.X)
            bgc, fgc, font = (C_SECONDARY, "white", ("微软雅黑", 10, "bold")) if r_idx == 0 else \
                             (("#f8fafc" if r_idx % 2 == 0 else C_CARD), C_TEXT_M, ("微软雅黑", 10))
            rc = tk.Frame(tf, bg=bgc); rc.pack(fill=tk.X)
            for cell, w in zip(row, widths):
                tk.Label(rc, text=cell, font=font, bg=bgc, fg=fgc, width=w, anchor="center").pack(side=tk.LEFT, padx=1, pady=3)

    def download_template(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="地址标准化模板.xlsx")
        if not path: return
        try:
            openpyxl, PatternFill, Font, Alignment, get_column_letter = _get_openpyxl()
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
            headers = ["序号", "省", "市", "区县", "地址1(必填)", "地址2(选填)", "GCJ-02经度", "GCJ-02纬度", "WGS-84经度", "WGS-84纬度", "标准地址", "完整地址", "省市区", "街道门牌号"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for r_idx, row in enumerate([
                [1, "四川省", "成都市", "", "西南电子科技大学", "", "", "", "", "", "", "", "", ""],
                [2, "北京市", "北京市", "", "天安门", "人民英雄纪念碑", "", "", "", "", "", "", "", ""],
                [3, "云南省", "曲靖市", "麒麟区", "曲靖市农村信用合作社联合社", "寥廓北路中段271-1号", "", "", "", "", "", "", "", ""]
            ], 2):
                for c_idx, v in enumerate(row, 1): ws.cell(row=r_idx, column=c_idx, value=v)
            for i, w in enumerate([8, 12, 12, 12, 25, 20, 15, 15, 15, 15, 40, 30, 20, 20], 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            wb.save(path)
            messagebox.showinfo("成功", f"模板已保存到：\n{path}")
        except Exception as e:
            messagebox.showerror("错误", f"保存失败：{e}")

    # ═══════════════════════════════════════════════════════
    # 单条查询
    # ═══════════════════════════════════════════════════════
    def _build_single(self):
        frame = self._make_page(self.content)
        self.pages["single"] = frame
        split = tk.Frame(frame, bg=C_BG)
        split.pack(fill=tk.BOTH, expand=True)
        split.columnconfigure(0, weight=35, minsize=400)
        split.columnconfigure(1, weight=65, minsize=700)
        split.rowconfigure(0, weight=1)

        left = tk.Frame(split, bg=C_BG)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 12))

        param_frame = Card(left, title=None, accent=None)
        param_frame.pack(fill=tk.X, pady=(0, 8))
        input_frame = tk.Frame(param_frame, bg=C_CARD)
        input_frame.pack(fill=tk.X, padx=14, pady=12)

        self.addr_frame = tk.Frame(input_frame, bg=C_CARD)
        self.addr_frame.pack(fill=tk.X)
        self.addr_label = tk.Label(self.addr_frame, text="* 地址（必填）", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT)
        self.addr_label.pack(anchor="w")
        self.addr_entry_frame = tk.Frame(self.addr_frame, bg=C_CARD)
        self.addr_entry_frame.pack(fill=tk.X, pady=(4, 0))
        self.single_addr = tk.Entry(self.addr_entry_frame, font=("微软雅黑", 11), relief=tk.SOLID, bd=1)
        self.single_addr.pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=3)
        self.single_addr.insert(0, "麒麟区雄业金都")
        tk.Button(self.addr_entry_frame, text="×", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT_LIGHT,
                  bd=0, cursor="hand2", command=lambda: self.single_addr.delete(0, tk.END)).pack(side=tk.RIGHT, padx=(4, 0))

        self.coord_label_frame = tk.Frame(input_frame, bg=C_CARD)
        tk.Label(self.coord_label_frame, text="经纬度", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT).pack(anchor="w")
        tk.Label(self.coord_label_frame, text="经度在前,纬度在后，英文逗号分割,如：103.791246,25.491221",
                font=("微软雅黑", 9), bg=C_CARD, fg=C_TEXT_LIGHT).pack(anchor="w")
        self.single_coord = tk.Entry(self.coord_label_frame, font=("微软雅黑", 11), relief=tk.SOLID, bd=1)
        self.single_coord.pack(fill=tk.X, pady=(4, 0), ipady=3)
        self.single_coord.insert(0, "103.791246,25.491221")
        sys_frame = tk.Frame(self.coord_label_frame, bg=C_CARD)
        sys_frame.pack(fill=tk.X, pady=(6, 0))
        tk.Label(sys_frame, text="坐标系：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.coord_sys = tk.StringVar(value="GCJ-02")
        for val, txt in [("GCJ-02", "GCJ-02"), ("WGS-84", "WGS-84"), ("BD-09", "BD-09")]:
            tk.Radiobutton(sys_frame, text=txt, variable=self.coord_sys, value=val,
                          font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(side=tk.LEFT, padx=(0, 12))

        btn_frame = tk.Frame(param_frame, bg=C_CARD)
        btn_frame.pack(fill=tk.X, padx=14, pady=(0, 12))
        SButton(btn_frame, "查询", self.do_single_query, "primary").pack(side=tk.LEFT, fill=tk.X, expand=True)
        mode_frame = tk.Frame(btn_frame, bg=C_CARD)
        mode_frame.pack(side=tk.RIGHT, padx=(10, 0))
        self.single_mode = tk.StringVar(value="address")
        tk.Radiobutton(mode_frame, text="根据地址", variable=self.single_mode, value="address",
                      font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD,
                      command=self._switch_single_mode).pack(side=tk.LEFT, padx=(0, 10))
        tk.Radiobutton(mode_frame, text="根据经纬度", variable=self.single_mode, value="coordinate",
                      font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD,
                      command=self._switch_single_mode).pack(side=tk.LEFT)

        result_frame = Card(left, title="标准化地址", accent=C_ACCENT)
        result_frame.pack(fill=tk.BOTH, expand=True)
        result_header = tk.Frame(result_frame, bg=C_CARD)
        result_header.pack(fill=tk.X, padx=14, pady=(10, 6))
        tk.Label(result_header, text="查询结果", font=("微软雅黑", 13, "bold"), bg=C_CARD, fg=C_ACCENT).pack(side=tk.LEFT)
        SButton(result_header, "📋 复制全部", self._copy_all_results, "secondary").pack(side=tk.RIGHT)
        tk.Frame(result_frame, bg=C_BORDER, height=1).pack(fill=tk.X, padx=14, pady=(0, 6))

        self.single_results = {}
        result_fields = [
            ("std_addr", "标准化地址"), ("gcj_lng", "GCJ-02经度"), ("gcj_lat", "GCJ-02纬度"),
            ("location", "坐标"), ("province", "省份"), ("city_r", "城市"),
            ("district", "区县"), ("township", "乡镇/街道"), ("street", "道路名称")
        ]
        for key, label in result_fields:
            row = tk.Frame(result_frame, bg=C_CARD)
            row.pack(fill=tk.X, padx=14, pady=3)
            tk.Label(row, text=label, font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT_M, width=11, anchor="e").pack(side=tk.LEFT)
            ent = tk.Entry(row, font=("微软雅黑", 10), relief=tk.SOLID, bd=1, state="readonly", bg="#f8fafc")
            ent.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=6, ipady=2)
            self.single_results[key] = ent

        self.single_error = tk.Label(result_frame, text="", font=("微软雅黑", 10), bg=C_CARD, fg=C_DANGER, wraplength=380)
        self.single_error.pack(fill=tk.X, padx=14, pady=(4, 2))

        tk.Label(result_frame, text="查询日志", font=("微软雅黑", 11, "bold"), bg=C_CARD, fg=C_TEXT).pack(anchor="w", padx=14, pady=(8, 4))
        tk.Frame(result_frame, bg=C_BORDER, height=1).pack(fill=tk.X, padx=14, pady=(0, 4))
        raw_frame = tk.Frame(result_frame, bg=C_CARD)
        raw_frame.pack(fill=tk.BOTH, expand=True, padx=14, pady=(0, 10))
        self.raw_text = scrolledtext.ScrolledText(raw_frame, font=("Consolas", 10), wrap=tk.WORD,
                                                   height=10, bg="#f8fafc", fg=C_TEXT, relief=tk.SOLID, bd=1)
        self.raw_text.pack(fill=tk.BOTH, expand=True)

        right = tk.Frame(split, bg=C_BG)
        right.grid(row=0, column=1, sticky="nsew")
        right.rowconfigure(0, weight=1)
        map_container = tk.Frame(right, bg="#e2e8f0", highlightbackground=C_BORDER, highlightthickness=1)
        map_container.pack(fill=tk.BOTH, expand=True)
        self.map_label = tk.Label(map_container, bg="#e2e8f0")
        self.map_label.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)
        self.map_label.image = None
        self.map_label.bind("<MouseWheel>", self._on_map_scroll)
        self.map_label.bind("<Button-4>", lambda e: self._on_map_scroll_dir(1))
        self.map_label.bind("<Button-5>", lambda e: self._on_map_scroll_dir(-1))
        self.zoom_level = 16
        self.map_lng = 116.397428
        self.map_lat = 39.90923

        MAP_BTN_BG = "#ffffff"; MAP_BTN_FG = "#475569"; MAP_BTN_ACTIVE = "#f1f5f9"
        zoom_frame = tk.Frame(map_container, bg=MAP_BTN_BG, highlightthickness=0)
        zoom_frame.place(x=12, y=12)
        for txt, cmd in [("+", self._zoom_in), ("−", self._zoom_out)]:
            tk.Button(zoom_frame, text=txt, font=("微软雅黑", 12, "bold"), bg=MAP_BTN_BG, fg=MAP_BTN_FG,
                      bd=0, width=2, height=1, relief=tk.FLAT, activebackground=MAP_BTN_ACTIVE,
                      command=cmd).pack(pady=(0, 1))
        pan_frame = tk.Frame(map_container, bg=MAP_BTN_BG, highlightthickness=0)
        pan_frame.place(relx=1.0, y=12, anchor="ne", x=-12)
        tk.Button(pan_frame, text="▲", font=("微软雅黑", 9), bg=MAP_BTN_BG, fg=MAP_BTN_FG,
                  bd=0, width=3, height=1, relief=tk.FLAT, activebackground=MAP_BTN_ACTIVE,
                  command=lambda: self._pan_map("up")).pack()
        mid_row = tk.Frame(pan_frame, bg=MAP_BTN_BG)
        mid_row.pack(fill=tk.X)
        tk.Button(mid_row, text="◀", font=("微软雅黑", 9), bg=MAP_BTN_BG, fg=MAP_BTN_FG,
                  bd=0, width=3, height=1, relief=tk.FLAT, activebackground=MAP_BTN_ACTIVE,
                  command=lambda: self._pan_map("left")).pack(side=tk.LEFT)
        tk.Label(mid_row, text="", font=("微软雅黑", 8), bg=MAP_BTN_BG, width=1).pack(side=tk.LEFT)
        tk.Button(mid_row, text="▶", font=("微软雅黑", 9), bg=MAP_BTN_BG, fg=MAP_BTN_FG,
                  bd=0, width=3, height=1, relief=tk.FLAT, activebackground=MAP_BTN_ACTIVE,
                  command=lambda: self._pan_map("right")).pack(side=tk.LEFT)
        tk.Button(pan_frame, text="▼", font=("微软雅黑", 9), bg=MAP_BTN_BG, fg=MAP_BTN_FG,
                  bd=0, width=3, height=1, relief=tk.FLAT, activebackground=MAP_BTN_ACTIVE,
                  command=lambda: self._pan_map("down")).pack()
        self.scale_label = tk.Label(map_container, text="", font=("微软雅黑", 9), bg=MAP_BTN_BG, fg=MAP_BTN_FG,
                                     bd=0, relief=tk.FLAT, padx=6, pady=3)
        self.scale_label.place(relx=1.0, rely=1.0, anchor="se", x=-12, y=-12)
        self.root.after(800, lambda: self._load_map(self.map_lng, self.map_lat))

    def _on_map_scroll(self, event):
        if event.delta > 0: self._zoom_in()
        else: self._zoom_out()

    def _on_map_scroll_dir(self, direction):
        if direction > 0: self._zoom_in()
        else: self._zoom_out()

    def _pan_map(self, direction):
        w = max(self.map_label.winfo_width(), 600)
        h = max(self.map_label.winfo_height(), 500)
        step_x, step_y = w // 4, h // 4
        dx = {"left": -step_x, "right": step_x}.get(direction, 0)
        dy = {"up": -step_y, "down": step_y}.get(direction, 0)
        zoom = self.zoom_level
        map_width = 256.0 * (2 ** zoom)
        lng_per_px = 360.0 / map_width
        lat_per_px = lng_per_px * math.cos(math.radians(self.map_lat))
        self.map_lng = max(-180, min(180, self.map_lng + dx * lng_per_px))
        self.map_lat = max(-85, min(85, self.map_lat - dy * lat_per_px))
        self._load_map(self.map_lng, self.map_lat)

    def _zoom_in(self):
        if self.zoom_level < 19: self.zoom_level += 1; self._load_map(self.map_lng, self.map_lat)

    def _zoom_out(self):
        if self.zoom_level > 3: self.zoom_level -= 1; self._load_map(self.map_lng, self.map_lat)

    def _update_scale(self):
        scale_map = {3: "1000 km", 4: "500 km", 5: "200 km", 6: "100 km", 7: "50 km", 8: "20 km", 9: "10 km",
                     10: "5 km", 11: "2 km", 12: "1 km", 13: "500 m", 14: "200 m", 15: "100 m", 16: "50 m",
                     17: "25 m", 18: "10 m", 19: "5 m"}
        self.scale_label.config(text=scale_map.get(self.zoom_level, f"zoom {self.zoom_level}"))

    def _switch_single_mode(self):
        mode = self.single_mode.get()
        if mode == "address":
            self.addr_label.config(text="* 地址（必填）")
            self.addr_frame.pack(fill=tk.X); self.coord_label_frame.pack_forget()
        else:
            self.addr_label.config(text="经纬度")
            self.addr_frame.pack_forget(); self.coord_label_frame.pack(fill=tk.X)
        self.clear_single()

    def clear_single(self):
        self.single_addr.delete(0, tk.END); self.single_addr.insert(0, "麒麟区雄业金都")
        self.single_coord.delete(0, tk.END); self.single_coord.insert(0, "103.791246,25.491221")
        for ent in self.single_results.values():
            ent.config(state="normal"); ent.delete(0, tk.END); ent.config(state="readonly")
        self.single_error.config(text="")
        self.raw_text.delete("1.0", tk.END)
        self.map_lng, self.map_lat = 116.397428, 39.90923
        self.zoom_level = 16
        self._load_map(self.map_lng, self.map_lat)
        self._last_result = {}

    def _copy_all_results(self):
        if not self._last_result:
            messagebox.showinfo("提示", "暂无查询结果可复制"); return
        lines = ["=" * 45, "地址坐标转换结果", "=" * 45]
        for key, label in [("std_addr", "标准化地址"), ("gcj_lng", "GCJ-02经度"), ("gcj_lat", "GCJ-02纬度"),
                           ("location", "坐标"), ("province", "省份"), ("city_r", "城市"),
                           ("district", "区县"), ("township", "乡镇/街道"), ("street", "道路名称")]:
            val = self._last_result.get(key, "")
            if val: lines.append(f"{label}: {val}")
        lines.append("=" * 45)
        self.root.clipboard_clear()
        self.root.clipboard_append("\n".join(lines))
        self.status_var.set("结果已复制到剪贴板")
        self.root.after(2000, lambda: self.status_var.set("就绪"))

    def do_single_query(self):
        mode = self.single_mode.get()
        if not self.config.get("gaode_key"):
            self.single_error.config(text="请先前往「API 设置」页面配置高德 API Key"); return
        self.geo = GeoCoder(self.config)
        self.status_var.set("正在查询...")
        self.single_error.config(text="")
        for ent in self.single_results.values():
            ent.config(state="normal"); ent.delete(0, tk.END); ent.config(state="readonly")
        self.raw_text.delete("1.0", tk.END)
        self._last_result = {}

        def task():
            ok, msg = self.geo.check_keys()
            if not ok:
                self.root.after(0, lambda: self.single_error.config(text=f"密钥错误：{msg}"))
                self.root.after(0, lambda: self.status_var.set("就绪")); return
            if mode == "address":
                addr = self.single_addr.get().strip()
                if not addr:
                    self.root.after(0, lambda: self.single_error.config(text="请输入地址"))
                    self.root.after(0, lambda: self.status_var.set("就绪")); return
                result, is_abnormal, err_msg, fatal, baidu_log = self.geo.resolve(addr, "", "", "")
                self.root.after(0, lambda: self._show_addr_result(result, is_abnormal, err_msg, fatal, baidu_log))
            else:
                coord_str = self.single_coord.get().strip()
                try:
                    parts = coord_str.split(",")
                    if len(parts) != 2: raise ValueError("格式错误")
                    lng, lat = float(parts[0]), float(parts[1])
                except Exception as e:
                    self.root.after(0, lambda: self.single_error.config(text=f"坐标格式错误：{e}"))
                    self.root.after(0, lambda: self.status_var.set("就绪")); return
                if self.coord_sys.get() == "WGS-84":
                    gcj_lng, gcj_lat = wgs84_to_gcj02(lng, lat)
                elif self.coord_sys.get() == "BD-09":
                    gcj_lng, gcj_lat = bd09_to_gcj02(lng, lat)
                else: gcj_lng, gcj_lat = lng, lat
                self._do_reverse_query(gcj_lng, gcj_lat)
        threading.Thread(target=task, daemon=True).start()

    def _show_addr_result(self, result, is_abnormal, err_msg, fatal, baidu_log):
        self.status_var.set("就绪")
        if fatal:
            self.single_error.config(text=f"查询失败（致命错误）：{err_msg}")
            self._build_query_log(None, is_abnormal, err_msg, fatal, baidu_log)
            return
        if not result:
            self.single_error.config(text=f"未找到结果：{err_msg}")
            self._build_query_log(None, is_abnormal, err_msg, fatal, baidu_log)
            return
        gcj_lng, gcj_lat = result["gcj_lng"], result["gcj_lat"]
        self.map_lng, self.map_lat = gcj_lng, gcj_lat
        fields = {"gcj_lng": str(gcj_lng), "gcj_lat": str(gcj_lat), "location": f"{gcj_lng},{gcj_lat}", "std_addr": result["std_addr"]}
        self._last_result.update(fields)
        for k, v in fields.items():
            if k in self.single_results:
                self.single_results[k].config(state="normal")
                self.single_results[k].delete(0, tk.END)
                self.single_results[k].insert(0, str(v))
                self.single_results[k].config(state="readonly")
        self._build_query_log(result, is_abnormal, err_msg, fatal, baidu_log)
        self._load_map(gcj_lng, gcj_lat)

        def regeo_task():
            regeo_result, err = self.geo.gaode_regeocode(gcj_lng, gcj_lat)
            if regeo_result:
                self.root.after(0, lambda: self._fill_regeo_result(regeo_result))
        threading.Thread(target=regeo_task, daemon=True).start()

    def _build_query_log(self, result, is_abnormal, err_msg, fatal, baidu_log):
        self.raw_text.delete("1.0", tk.END)
        lines = []
        lines.append("=" * 50)
        lines.append("查询日志")
        lines.append("=" * 50)
        if fatal:
            lines.append(f"[致命错误] {err_msg}")
            lines.append("=" * 50)
            self.raw_text.insert(tk.END, "\n".join(lines))
            return
        if not result:
            lines.append(f"[查询失败] {err_msg or '未找到结果'}")
            lines.append("=" * 50)
            self.raw_text.insert(tk.END, "\n".join(lines))
            return
        lines.append(f"[高德查询] 地址: {result.get('std_addr', '')}")
        lines.append(f"[高德查询] 精度: {result.get('level', '未知')}")
        lines.append(f"[高德查询] 坐标: {result.get('gcj_lng', '')}, {result.get('gcj_lat', '')}")
        if result.get('is_drift'):
            lines.append("[高德查询] 警告: 疑似地址漂移")
        if baidu_log:
            lines.append("-" * 50)
            for log_line in baidu_log.strip().split("\n"):
                if log_line.strip():
                    lines.append(log_line.strip())
        lines.append("-" * 50)
        source = result.get('source', '高德')
        level = result.get('level', '未知')
        if is_abnormal:
            lines.append(f"[最终结果] 状态: 异常 (精度不足: {level})")
        else:
            lines.append(f"[最终结果] 状态: 成功")
        lines.append(f"[最终结果] 来源: {source}")
        lines.append(f"[最终结果] 精度: {level}")
        lines.append(f"[最终结果] 标准化地址: {result.get('std_addr', '')}")
        lines.append("=" * 50)
        self.raw_text.insert(tk.END, "\n".join(lines))

    def _fill_regeo_result(self, regeo):
        for k, v in [("province", regeo.get("province", "")), ("city_r", regeo.get("city", "")),
                     ("district", regeo.get("district", "")), ("township", regeo.get("township", "")),
                     ("street", regeo.get("street", ""))]:
            if k in self.single_results:
                self.single_results[k].config(state="normal")
                self.single_results[k].delete(0, tk.END)
                self.single_results[k].insert(0, str(v))
                self.single_results[k].config(state="readonly")
                self._last_result[k] = str(v)

    def _do_reverse_query(self, gcj_lng, gcj_lat):
        self.map_lng, self.map_lat = gcj_lng, gcj_lat
        def task():
            regeo_result, err = self.geo.gaode_regeocode(gcj_lng, gcj_lat)
            self.root.after(0, lambda: self._show_regeo_result(gcj_lng, gcj_lat, regeo_result, err))
        threading.Thread(target=task, daemon=True).start()

    def _show_regeo_result(self, gcj_lng, gcj_lat, regeo_result, err):
        self.status_var.set("就绪")
        if not regeo_result: self.single_error.config(text=f"逆地理编码失败：{err}"); return
        fields = {"gcj_lng": str(gcj_lng), "gcj_lat": str(gcj_lat), "location": f"{gcj_lng},{gcj_lat}",
                  "province": regeo_result.get("province", ""), "city_r": regeo_result.get("city", ""),
                  "district": regeo_result.get("district", ""), "township": regeo_result.get("township", ""),
                  "street": regeo_result.get("street", ""), "std_addr": regeo_result.get("full", "")}
        self._last_result.update(fields)
        for k, v in fields.items():
            if k in self.single_results:
                self.single_results[k].config(state="normal")
                self.single_results[k].delete(0, tk.END)
                self.single_results[k].insert(0, str(v))
                self.single_results[k].config(state="readonly")
        self.raw_text.delete("1.0", tk.END)
        self.raw_text.insert(tk.END, json.dumps(regeo_result, ensure_ascii=False, indent=2))
        self._load_map(gcj_lng, gcj_lat)

    def _load_map(self, lng, lat):
        def task():
            key = self.config.get("gaode_key", "")
            zoom = self.zoom_level
            Image, ImageTk = _get_pil()
            raw_w = self.map_label.winfo_width()
            raw_h = self.map_label.winfo_height()
            if raw_w < 100 or raw_h < 100:
                raw_w = max(self.root.winfo_width() - 640, 600)
                raw_h = max(self.root.winfo_height() - 120, 500)
            w = max(raw_w, 600)
            h = max(raw_h, 500)
            if key:
                try:
                    url = (f"{GAODE_STATIC}?location={lng},{lat}&zoom={zoom}"
                           f"&size={w}*{h}&markers=mid,0xFF0000,{lng},{lat}&key={key}")
                    ok, data = self._fetch_image(url)
                    if ok:
                        img = Image.open(io.BytesIO(data))
                        self.root.after(0, lambda img=img: self._set_map_image(img))
                        self.root.after(0, self._update_scale); return
                except: pass
            baidu_ak = self.config.get("baidu_ak", "")
            if baidu_ak:
                try:
                    bd_lng, bd_lat = gcj02_to_bd09(lng, lat)
                    url = (f"{BAIDU_STATIC}?ak={baidu_ak}&center={bd_lng},{bd_lat}"
                           f"&width={w}&height={h}&zoom={zoom}&markers={bd_lng},{bd_lat}")
                    ok, data = self._fetch_image(url)
                    if ok:
                        img = Image.open(io.BytesIO(data))
                        self.root.after(0, lambda img=img: self._set_map_image(img))
                        self.root.after(0, self._update_scale); return
                except: pass
            err_msg = "地图加载失败\n"
            if not key: err_msg += "未配置高德 API Key\n"
            else: err_msg += "高德静态地图服务异常（权限不足或服务端错误 20003）\n"
            if not baidu_ak: err_msg += "未配置百度 AK，无法使用备用地图源\n"
            err_msg += "\n可能原因：\n1. 高德Key类型不正确（需Web服务Key）\n2. 静态地图API未开通或配额超限\n3. 数字签名已开启但未传入签名参数\n4. 网络连接问题"
            self.root.after(0, lambda: self.map_label.config(text=err_msg))
        threading.Thread(target=task, daemon=True).start()

    def _fetch_image(self, url):
        ctx = ssl.create_default_context()
        ctx.check_hostname, ctx.verify_mode = False, ssl.CERT_NONE
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15, context=ctx) as resp:
            content_type = resp.headers.get('Content-Type', '')
            data = resp.read()
            if b'application/json' in content_type.encode() or data.startswith(b'{'):
                return False, None
            return True, data

    def _set_map_image(self, img):
        try:
            Image, ImageTk = _get_pil()
            w = max(self.map_label.winfo_width(), 600)
            h = max(self.map_label.winfo_height(), 500)
            img = img.resize((w, h), Image.LANCZOS)
            photo = ImageTk.PhotoImage(img)
            self.map_label.config(image=photo, text="")
            self.map_label.image = photo
        except Exception as e:
            self.map_label.config(text=f"图片处理失败: {str(e)[:80]}")

    # ═══════════════════════════════════════════════════════
    # 批量编码（通用构建）
    # ═══════════════════════════════════════════════════════
    def _build_batch(self, direction):
        """构建批量编码页面。direction: 'fwd' 或 'rev'"""
        is_fwd = direction == "fwd"
        page_key = f"batch_{direction}"
        frame = self._make_page(self.content)
        self.pages[page_key] = frame
        title = "批量正向地理编码（地址 -> 坐标）" if is_fwd else "批量逆向地理编码（坐标 -> 地址）"
        tk.Label(frame, text=title, font=("微软雅黑", 18, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor="w", pady=(0, 10))

        split = tk.Frame(frame, bg=C_BG)
        split.pack(fill=tk.BOTH, expand=True)
        split.columnconfigure(0, weight=0, minsize=360)
        split.columnconfigure(1, weight=1, minsize=650)
        split.rowconfigure(0, weight=1)

        # 左侧：带垂直滚动的 Canvas 面板，防止内容被截断
        left_outer = tk.Frame(split, bg=C_CARD, highlightbackground=C_BORDER, highlightthickness=1)
        left_outer.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        left_outer.pack_propagate(False)
        left_outer.config(width=380)
        left_outer.grid_rowconfigure(0, weight=1)
        left_outer.grid_columnconfigure(0, weight=1)

        canvas = tk.Canvas(left_outer, bg=C_CARD, highlightthickness=0)
        canvas.grid(row=0, column=0, sticky="nsew")
        vsb = ttk.Scrollbar(left_outer, orient="vertical", command=canvas.yview)
        vsb.grid(row=0, column=1, sticky="ns")
        canvas.configure(yscrollcommand=vsb.set)

        left_inner = tk.Frame(canvas, bg=C_CARD)
        canvas_window = canvas.create_window((0, 0), window=left_inner, anchor="nw", width=360)
        def _on_canvas_configure(event):
            canvas.itemconfig(canvas_window, width=event.width)
        canvas.bind("<Configure>", _on_canvas_configure)
        def _on_frame_configure(event):
            canvas.configure(scrollregion=canvas.bbox("all"))
        left_inner.bind("<Configure>", _on_frame_configure)
        # 鼠标滚轮绑定
        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1*(event.delta/120)), "units")
        canvas.bind("<MouseWheel>", _on_mousewheel)
        left_inner.bind("<MouseWheel>", _on_mousewheel)

        left_inner = tk.Frame(left_inner, bg=C_CARD)
        left_inner.pack(fill=tk.BOTH, expand=True, padx=12, pady=10)

        tk.Label(left_inner, text="▎第一步：导入文件", font=("微软雅黑", 12, "bold"), bg=C_CARD, fg=C_ACCENT).pack(anchor="w", pady=(0, 6))
        path_var = tk.StringVar()
        path_entry = tk.Entry(left_inner, textvariable=path_var, font=("微软雅黑", 10), relief=tk.SOLID, bd=1, state="readonly", bg="#f8fafc")
        path_entry.pack(fill=tk.X, pady=(0, 6))
        setattr(self, f"{direction}_path_var", path_var)
        import_btn = tk.Button(left_inner, text="📂 导入 Excel 文件", font=("微软雅黑", 11, "bold"),
                               bg=C_ACCENT, fg="#fff", bd=0, padx=20, pady=10, cursor="hand2",
                               relief=tk.FLAT, command=lambda: self._import_batch(direction))
        import_btn.pack(fill=tk.X, pady=(0, 8))
        import_btn.bind("<Enter>", lambda e: import_btn.config(bg="#b45309"))
        import_btn.bind("<Leave>", lambda e: import_btn.config(bg=C_ACCENT))

        tk.Frame(left_inner, bg=C_BORDER, height=1).pack(fill=tk.X, pady=(0, 6))

        tk.Label(left_inner, text="▎第二步：指定数据列", font=("微软雅黑", 12, "bold"), bg=C_CARD, fg=C_ACCENT).pack(anchor="w", pady=(0, 5))
        col_frame = tk.Frame(left_inner, bg=C_CARD)
        col_frame.pack(fill=tk.X, pady=(0, 2))

        if is_fwd:
            labels_combos = [
                ("* 地址1列（必填）", "fwd_addr_col"),
                ("* 地址2列（必填，备用/精度补充）", "fwd_addr2_col"),
                ("  城市列（可选，提升精度）", "fwd_city_col")
            ]
        else:
            labels_combos = [
                ("* GCJ-02 经度列", "rev_gcj_lng_col"),
                ("* GCJ-02 纬度列", "rev_gcj_lat_col"),
                ("  WGS-84 经度列（备选）", "rev_wgs_lng_col"),
                ("  WGS-84 纬度列（备选）", "rev_wgs_lat_col")
            ]
        for lbl_text, attr_name in labels_combos:
            fg = C_TEXT if "*" in lbl_text else C_TEXT_LIGHT
            tk.Label(col_frame, text=lbl_text, font=("微软雅黑", 10), bg=C_CARD, fg=fg).pack(anchor="w")
            cb = ttk.Combobox(col_frame, values=[], font=("微软雅黑", 10), width=30, state="readonly")
            cb.pack(fill=tk.X, pady=(2, 4))
            setattr(self, attr_name, cb)

        tk.Frame(left_inner, bg=C_BORDER, height=1).pack(fill=tk.X, pady=(0, 6))

        tk.Label(left_inner, text="▎第三步：处理参数", font=("微软雅黑", 12, "bold"), bg=C_CARD, fg=C_ACCENT).pack(anchor="w", pady=(0, 5))
        param_grid = tk.Frame(left_inner, bg=C_CARD)
        param_grid.pack(fill=tk.X, pady=(0, 4))
        tk.Label(param_grid, text="起始行：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).grid(row=0, column=0, sticky="w")
        start_spin = tk.Spinbox(param_grid, from_=2, to=100000, width=10, font=("微软雅黑", 10), relief=tk.SOLID, bd=1)
        start_spin.grid(row=0, column=1, sticky="w", padx=(8, 0), pady=4)
        start_spin.delete(0, tk.END); start_spin.insert(0, "2")
        setattr(self, f"{direction}_start", start_spin)

        skip_var = tk.BooleanVar(value=True)
        tk.Checkbutton(left_inner, text="跳过已有结果", variable=skip_var,
                       font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(anchor="w", pady=2)
        setattr(self, f"{direction}_skip", skip_var)

        if is_fwd:
            ab_var = tk.BooleanVar(value=False)
            tk.Checkbutton(left_inner, text="只重新处理异常行", variable=ab_var,
                           font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(anchor="w", pady=2)
            setattr(self, "fwd_abnormal", ab_var)

        auto_var = tk.BooleanVar(value=self.config.get("auto_open_excel", False))
        tk.Checkbutton(left_inner, text="完成后自动打开文件", variable=auto_var,
                       font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(anchor="w", pady=2)
        setattr(self, f"{direction}_auto_open", auto_var)

        tk.Frame(left_inner, bg=C_BORDER, height=1).pack(fill=tk.X, pady=(8, 6))

        tk.Label(left_inner, text="▎第四步：开始转换", font=("微软雅黑", 12, "bold"), bg=C_CARD, fg=C_ACCENT).pack(anchor="w", pady=(0, 6))
        btn_row1 = tk.Frame(left_inner, bg=C_CARD)
        btn_row1.pack(fill=tk.X, pady=(0, 5))
        start_btn = tk.Button(btn_row1, text="▶ 开始处理", font=("微软雅黑", 11, "bold"),
                              bg=C_ACCENT, fg=C_PRIMARY, bd=0, padx=20, pady=10,
                              cursor="hand2", relief=tk.FLAT, command=lambda: self._start_batch(direction))
        start_btn.pack(side=tk.LEFT, fill=tk.X, expand=True)
        start_btn.bind("<Enter>", lambda e: start_btn.config(bg="#b45309"))
        start_btn.bind("<Leave>", lambda e: start_btn.config(bg=C_ACCENT))
        setattr(self, f"{direction}_start_btn", start_btn)

        stop_btn = tk.Button(btn_row1, text="⏹ 停止", font=("微软雅黑", 11, "bold"),
                             bg=C_DANGER, fg="white", bd=0, padx=20, pady=10,
                             cursor="hand2", relief=tk.FLAT, state=tk.DISABLED, command=self.stop_batch)
        stop_btn.pack(side=tk.LEFT, padx=(8, 0), fill=tk.X, expand=True)
        setattr(self, f"{direction}_stop_btn", stop_btn)

        btn_row2 = tk.Frame(left_inner, bg=C_CARD)
        btn_row2.pack(fill=tk.X, pady=(0, 5))
        export_btn = tk.Button(btn_row2, text="📤 导出异常行", font=("微软雅黑", 10, "bold"),
                               bg=C_WARNING, fg="#fff", bd=0, padx=20, pady=10,
                               cursor="hand2", relief=tk.FLAT, command=lambda: self._export_batch_abnormal(direction))
        export_btn.pack(fill=tk.X)
        export_btn.bind("<Enter>", lambda e: export_btn.config(bg="#d97706"))
        export_btn.bind("<Leave>", lambda e: export_btn.config(bg=C_WARNING))

        prog_frame = tk.Frame(left_inner, bg=C_CARD)
        prog_frame.pack(fill=tk.X, pady=(0, 3))
        prog = ttk.Progressbar(prog_frame, orient=tk.HORIZONTAL, mode="determinate")
        prog.pack(fill=tk.X)
        prog_lbl = tk.Label(prog_frame, text="就绪", font=("微软雅黑", 9), bg=C_CARD, fg=C_TEXT_LIGHT)
        prog_lbl.pack(anchor="w", pady=(4, 0))
        setattr(self, f"{direction}_prog", prog)
        setattr(self, f"{direction}_prog_lbl", prog_lbl)

        right = tk.Frame(split, bg=C_BG)
        right.grid(row=0, column=1, sticky="nsew")
        right.rowconfigure(0, weight=1)
        right.columnconfigure(0, weight=1)
        notebook = ttk.Notebook(right)
        notebook.pack(fill=tk.BOTH, expand=True)
        setattr(self, f"{direction}_notebook", notebook)

        preview = PreviewManager(right, notebook, "📋 数据预览")
        setattr(self, f"{direction}_preview", preview)

        log_frame = tk.Frame(notebook, bg=C_CARD)
        notebook.add(log_frame, text="  📊 进度预览  ")
        log_widget = scrolledtext.ScrolledText(log_frame, font=("Consolas", 10), wrap=tk.WORD,
                                                height=14, bg="#f8fafc", fg=C_TEXT,
                                                relief=tk.SOLID, bd=1, padx=10, pady=10)
        log_widget.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self._config_log_tags(log_widget)
        setattr(self, f"{direction}_log", log_widget)

    def _import_batch(self, direction):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xls")])
        if not path: return
        getattr(self, f"{direction}_path_var").set(path)
        self._load_batch_preview(path, direction)

    def _load_batch_preview(self, path, direction):
        try:
            openpyxl = _get_openpyxl()[0]
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if "Sheet1" not in wb.sheetnames:
                messagebox.showwarning("提示", "Excel 中未找到 Sheet1")
                wb.close(); return
            ws = wb["Sheet1"]
            headers = [str(cell.value).strip() if cell.value else f"列{cell.column}" for cell in ws[1]]
            data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(data) >= 200: break
                row_dict = {headers[i] if i < len(headers) else f"列{i+1}": (str(v) if v is not None else "") for i, v in enumerate(row)}
                data.append(row_dict)
            wb.close()
            getattr(self, f"{direction}_preview").load_data(headers, data)
            self._auto_select_columns(headers, direction)
            getattr(self, f"{direction}_notebook").select(0)
            self.status_var.set(f"已加载文件：{os.path.basename(path)}")
        except Exception as e:
            messagebox.showerror("错误", f"读取文件失败：{e}")

    def _auto_select_columns(self, headers, direction):
        if direction == "fwd":
            addr_match = addr2_match = city_match = None
            for h in headers:
                hl = h.lower()
                if addr_match is None and any(k in hl for k in ["地址1", "地址一", "addr1", "address1", "详细地址"]):
                    addr_match = h
                if addr2_match is None and any(k in hl for k in ["地址2", "地址二", "addr2", "address2", "备用地址"]):
                    addr2_match = h
                if city_match is None and any(k in hl for k in ["市", "city", "城市"]):
                    city_match = h
            if addr_match: self.fwd_addr_col.set(addr_match)
            if addr2_match: self.fwd_addr2_col.set(addr2_match)
            if city_match: self.fwd_city_col.set(city_match)
        else:
            gcj_lng_match = gcj_lat_match = wgs_lng_match = wgs_lat_match = None
            for h in headers:
                hl = h.lower()
                if gcj_lng_match is None and any(k in hl for k in ["gcj-02经度", "gcj02经度", "gcj经度", "火星经度"]):
                    gcj_lng_match = h
                if gcj_lat_match is None and any(k in hl for k in ["gcj-02纬度", "gcj02纬度", "gcj纬度", "火星纬度"]):
                    gcj_lat_match = h
                if wgs_lng_match is None and any(k in hl for k in ["wgs-84经度", "wgs84经度", "wgs经度", "84经度"]):
                    wgs_lng_match = h
                if wgs_lat_match is None and any(k in hl for k in ["wgs-84纬度", "wgs84纬度", "wgs纬度", "84纬度"]):
                    wgs_lat_match = h
            if not gcj_lng_match:
                for h in headers:
                    if any(k in h.lower() for k in ["经度", "lng", "lon", "x"]):
                        gcj_lng_match = h; break
            if not gcj_lat_match:
                for h in headers:
                    if any(k in h.lower() for k in ["纬度", "lat", "y"]):
                        gcj_lat_match = h; break
            if gcj_lng_match: self.rev_gcj_lng_col.set(gcj_lng_match)
            if gcj_lat_match: self.rev_gcj_lat_col.set(gcj_lat_match)
            if wgs_lng_match: self.rev_wgs_lng_col.set(wgs_lng_match)
            if wgs_lat_match: self.rev_wgs_lat_col.set(wgs_lat_match)

    def _export_batch_abnormal(self, direction):
        path = getattr(self, f"{direction}_path_var").get()
        if not path or not os.path.exists(path):
            messagebox.showwarning("提示", "请先选择有效的 Excel 文件"); return
        try: start_row = int(getattr(self, f"{direction}_start").get())
        except: start_row = 2
        output = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                               initialfile=f"异常行_{'正向' if direction == 'fwd' else '逆向'}编码.xlsx")
        if not output: return
        def task():
            try:
                proc = ExcelProcessor(path)
                ok, msg = proc.export_abnormal_rows(output, start_row)
                self.root.after(0, lambda: messagebox.showinfo("导出结果", msg) if ok else messagebox.showwarning("导出失败", msg))
            except Exception as e:
                self.root.after(0, lambda: messagebox.showerror("错误", f"导出异常：{e}"))
        threading.Thread(target=task, daemon=True).start()

    def _start_batch(self, direction):
        is_fwd = direction == "fwd"
        path = getattr(self, f"{direction}_path_var").get()
        if not path or not os.path.exists(path):
            messagebox.showwarning("提示", "请选择有效的 Excel 文件"); return
        if not self.config.get("gaode_key"):
            messagebox.showwarning("提示", "请先配置高德 API Key"); self.show_page("settings"); return
        try: start_row = int(getattr(self, f"{direction}_start").get())
        except: start_row = 2
        self.running, self.stop_flag = True, False
        getattr(self, f"{direction}_start_btn").config(state=tk.DISABLED)
        getattr(self, f"{direction}_stop_btn").config(state=tk.NORMAL)
        getattr(self, f"{direction}_log").delete("1.0", tk.END)
        auto_open = getattr(self, f"{direction}_auto_open").get()
        self.config["auto_open_excel"] = auto_open
        save_config(self.config)

        def task():
            try:
                proc = ExcelProcessor(path)
                headers = [str(proc.ws.cell(row=1, column=c).value or "").strip()
                          for c in range(1, proc.ws.max_column + 1)]
                if is_fwd:
                    addr_idx = self._get_col_idx(headers, self.fwd_addr_col.get(), COLS["addr1"])
                    addr2_idx = self._get_col_idx(headers, self.fwd_addr2_col.get(), COLS["addr2"])
                    city_idx = self._get_col_idx(headers, self.fwd_city_col.get(), COLS["city"])
                    if not self.fwd_addr_col.get() or not self.fwd_addr2_col.get():
                        self.root.after(0, lambda: self._log_batch(direction, "[警告] 地址1列和地址2列必须同时指定\n", "warning"))
                    last_row = proc.get_last_row(addr_idx)
                else:
                    gcj_lng_idx = self._get_col_idx(headers, self.rev_gcj_lng_col.get(), COLS["gcj_lng"])
                    gcj_lat_idx = self._get_col_idx(headers, self.rev_gcj_lat_col.get(), COLS["gcj_lat"])
                    wgs_lng_idx = self._get_col_idx(headers, self.rev_wgs_lng_col.get(), 0) if self.rev_wgs_lng_col.get() else 0
                    wgs_lat_idx = self._get_col_idx(headers, self.rev_wgs_lat_col.get(), 0) if self.rev_wgs_lat_col.get() else 0
                    last_row = max(proc.get_last_row(gcj_lng_idx), proc.get_last_row(gcj_lat_idx),
                                   proc.get_last_row(wgs_lng_idx) if wgs_lng_idx > 0 else 1,
                                   proc.get_last_row(wgs_lat_idx) if wgs_lat_idx > 0 else 1)
                    if last_row < 2: last_row = proc.ws.max_row
                total = last_row - start_row + 1
                if total <= 0:
                    self.root.after(0, lambda: self._log_batch(direction, "提示：没有找到需要处理的数据\n", "info")); return
                ok, msg = self.batch_processor.init(self.config)
                if not ok:
                    self.root.after(0, lambda: self._log_batch(direction, f"[密钥错误] {msg}\n", "error")); return

                def log_cb(msg, tag="normal"): self.root.after(0, lambda: self._log_batch(direction, msg, tag))
                def prog_cb(val, total_val):
                    self.root.after(0, lambda: getattr(self, f"{direction}_prog_lbl").config(text=f"{val} / {total_val}"))
                    self.root.after(0, lambda: getattr(self, f"{direction}_prog").config(value=val))

                self.root.after(0, lambda: getattr(self, f"{direction}_prog").config(maximum=total, value=0))
                if is_fwd:
                    stats = self.batch_processor.process_forward(
                        proc, start_row, last_row, addr_idx, addr2_idx, city_idx,
                        self.fwd_skip.get(), self.fwd_abnormal.get(), log_cb, prog_cb
                    )
                    self.root.after(0, lambda: self._log_batch(direction,
                        f"\n{'='*45}\n处理完成统计\n{'='*45}\n"
                        f"总行数：{stats['total']}\n成功：{stats['success']}\n使用地址2：{stats['addr2_used']}\n"
                        f"百度优化提升：{stats['baidu_up']}\n疑似漂移：{stats['drift_cnt']}\n失败：{stats['fail']}\n"
                        f"跳过：{stats['skip']}\n高德调用：{stats['gaode_count']}\n"
                        f"百度调用：{stats['baidu_count']}\n{'='*45}\n", "info"))
                else:
                    stats = self.batch_processor.process_reverse(
                        proc, start_row, last_row, gcj_lng_idx, gcj_lat_idx,
                        wgs_lng_idx, wgs_lat_idx, self.rev_skip.get(), log_cb, prog_cb
                    )
                    if stats.get("global_error"):
                        self.root.after(0, lambda: self._log_batch(direction, "[全局错误] 触发全局错误，已终止处理！\n", "error"))
                    else:
                        self.root.after(0, lambda: self._log_batch(direction,
                            f"\n{'='*45}\n处理完成统计\n{'='*45}\n"
                            f"总行数：{stats['total']}\n成功：{stats['success']}\n失败：{stats['fail']}\n"
                            f"跳过：{stats['skip']}\n高德调用：{stats['gaode_count']}\n{'='*45}\n", "info"))
                proc.save()
                if auto_open and os.path.exists(path): self._open_file(path)
            except Exception as e:
                self.root.after(0, lambda: self._log_batch(direction, f"[处理异常] {str(e)}\n", "error"))
            finally:
                self.running = False
                self.root.after(0, lambda: (getattr(self, f"{direction}_start_btn").config(state=tk.NORMAL),
                                             getattr(self, f"{direction}_stop_btn").config(state=tk.DISABLED),
                                             self.status_var.set("就绪")))
        threading.Thread(target=task, daemon=True).start()

    def stop_batch(self):
        self.stop_flag = True
        self.batch_processor.stop()
        self.status_var.set("正在停止...")

    def _open_file(self, path):
        try:
            sys_platform = platform.system()
            if sys_platform == "Windows": os.startfile(path)
            elif sys_platform == "Darwin": subprocess.run(["open", path], check=False)
            else: subprocess.run(["xdg-open", path], check=False)
        except Exception as e:
            print(f"自动打开文件失败: {e}")

    # ═══════════════════════════════════════════════════════
    # 坐标转换
    # ═══════════════════════════════════════════════════════
    def _build_converter(self):
        frame = self._make_page(self.content)
        self.pages["converter"] = frame
        tk.Label(frame, text="坐标转换工具", font=("微软雅黑", 20, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor="w", pady=(0, 6))
        tk.Label(frame, text="GCJ-02 / WGS-84 / BD-09 三系坐标互转，支持单点转换与批量Excel转换",
                 font=("微软雅黑", 11), bg=C_BG, fg=C_TEXT_LIGHT).pack(anchor="w", pady=(0, 16))

        single_card = Card(frame, title="▎单点转换", accent=C_INFO)
        single_card.pack(fill=tk.X, pady=(0, 6))
        sf = tk.Frame(single_card, bg=C_CARD)
        sf.pack(fill=tk.X, padx=14, pady=8)
        row_frame = tk.Frame(sf, bg=C_CARD)
        row_frame.pack(fill=tk.X, pady=(0, 6))
        tk.Label(row_frame, text="输入坐标：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_input = tk.Entry(row_frame, font=("微软雅黑", 11), width=30, relief=tk.SOLID, bd=1)
        self.cv_input.pack(side=tk.LEFT, padx=10, ipady=2)
        self.cv_input.insert(0, "103.791246,25.491221")
        tk.Label(row_frame, text="转换方式：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT, padx=(15, 0))
        self.cv_trans_single = ttk.Combobox(row_frame, values=[
            "GCJ-02 → WGS-84", "GCJ-02 → BD-09", "WGS-84 → GCJ-02",
            "WGS-84 → BD-09", "BD-09 → GCJ-02", "BD-09 → WGS-84"
        ], font=("微软雅黑", 10), width=18, state="readonly")
        self.cv_trans_single.set("GCJ-02 → WGS-84")
        self.cv_trans_single.pack(side=tk.LEFT, padx=10)
        SButton(row_frame, "  转换", self._do_single_convert, "primary").pack(side=tk.LEFT, padx=10)
        tk.Label(row_frame, text="格式：经度,纬度", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT_LIGHT).pack(side=tk.LEFT, padx=(5, 0))

        out_frame = tk.Frame(sf, bg=C_CARD)
        out_frame.pack(fill=tk.X, pady=(0, 4))
        tk.Label(out_frame, text="转换结果：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_output = tk.Entry(out_frame, font=("微软雅黑", 11), width=38, relief=tk.SOLID, bd=1, state="readonly", bg="#f8fafc")
        self.cv_output.pack(side=tk.LEFT, padx=10, ipady=2)
        SButton(out_frame, "📋 复制", self._copy_convert_result, "secondary").pack(side=tk.LEFT)
        self.cv_error = tk.Label(sf, text="", font=("微软雅黑", 10), bg=C_CARD, fg=C_DANGER)
        self.cv_error.pack(anchor="w", pady=(6, 0))

        batch_card = Card(frame, title="▎批量Excel转换", accent=C_SUCCESS)
        batch_card.pack(fill=tk.BOTH, expand=True, pady=(6, 0))
        bf = tk.Frame(batch_card, bg=C_CARD)
        bf.pack(fill=tk.X, padx=14, pady=8)

        file_row = tk.Frame(bf, bg=C_CARD)
        file_row.pack(fill=tk.X, pady=(0, 6))
        tk.Label(file_row, text="Excel 文件：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_batch_path = tk.StringVar()
        self.cv_batch_entry = tk.Entry(file_row, textvariable=self.cv_batch_path, font=("微软雅黑", 10), width=50, relief=tk.SOLID, bd=1)
        self.cv_batch_entry.pack(side=tk.LEFT, padx=10, fill=tk.X, expand=True, ipady=2)
        tk.Button(file_row, text="浏览...", font=("微软雅黑", 10), bg=C_BORDER, fg=C_TEXT, bd=0, padx=15, pady=5,
                 command=self._select_cv_batch_file).pack(side=tk.LEFT)
        tk.Button(file_row, text="读取列", font=("微软雅黑", 10), bg=C_INFO, fg="white", bd=0, padx=15, pady=5,
                 command=self._load_cv_columns).pack(side=tk.LEFT, padx=6)

        self.cv_col_frame = tk.Frame(bf, bg=C_CARD)
        self.cv_col_frame.pack(fill=tk.X, pady=(0, 10))
        self.cv_col_frame.pack_forget()

        src_group = tk.LabelFrame(self.cv_col_frame, text=" 源数据与转换方式 ", font=("微软雅黑", 10, "bold"), bg=C_CARD, fg=C_TEXT, bd=1, relief=tk.SOLID)
        src_group.pack(fill=tk.X, pady=6)
        src_inner = tk.Frame(src_group, bg=C_CARD)
        src_inner.pack(fill=tk.X, padx=10, pady=8)
        tk.Label(src_inner, text="源经度列：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_lng_col = ttk.Combobox(src_inner, values=[], font=("微软雅黑", 10), width=16, state="readonly")
        self.cv_lng_col.pack(side=tk.LEFT, padx=(0, 14))
        tk.Label(src_inner, text="源纬度列：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_lat_col = ttk.Combobox(src_inner, values=[], font=("微软雅黑", 10), width=16, state="readonly")
        self.cv_lat_col.pack(side=tk.LEFT, padx=(0, 14))
        tk.Label(src_inner, text="转换方式：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_trans_mode = ttk.Combobox(src_inner, values=[
            "GCJ-02 → WGS-84", "GCJ-02 → BD-09", "WGS-84 → GCJ-02",
            "WGS-84 → BD-09", "BD-09 → GCJ-02", "BD-09 → WGS-84"
        ], font=("微软雅黑", 10), width=18, state="readonly")
        self.cv_trans_mode.set("GCJ-02 → WGS-84")
        self.cv_trans_mode.pack(side=tk.LEFT)
        self.cv_trans_mode.bind("<<ComboboxSelected>>", self._on_trans_mode_changed)

        out_group = tk.LabelFrame(self.cv_col_frame, text=" 输出设置（选择Excel中目标列） ", font=("微软雅黑", 10, "bold"), bg=C_CARD, fg=C_TEXT, bd=1, relief=tk.SOLID)
        out_group.pack(fill=tk.X, pady=6)
        out_inner = tk.Frame(out_group, bg=C_CARD)
        out_inner.pack(fill=tk.X, padx=10, pady=8)
        tk.Label(out_inner, text="输出经度列：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_out_lng_col = ttk.Combobox(out_inner, values=[], font=("微软雅黑", 10), width=20, state="readonly")
        self.cv_out_lng_col.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(out_inner, text="输出纬度列：", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT)
        self.cv_out_lat_col = ttk.Combobox(out_inner, values=[], font=("微软雅黑", 10), width=20, state="readonly")
        self.cv_out_lat_col.pack(side=tk.LEFT)
        self.cv_skip = tk.BooleanVar(value=True)
        tk.Checkbutton(out_inner, text="跳过已有结果", variable=self.cv_skip,
                       font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(side=tk.LEFT, padx=(20, 10))
        self.cv_auto_open = tk.BooleanVar(value=self.config.get("auto_open_excel", False))
        tk.Checkbutton(out_inner, text="完成后自动打开文件", variable=self.cv_auto_open,
                       font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(side=tk.LEFT)

        btn_prog_row = tk.Frame(self.cv_col_frame, bg=C_CARD)
        btn_prog_row.pack(fill=tk.X, pady=4)
        self.cv_start_btn = SButton(btn_prog_row, "  开始批量转换", self._start_batch_convert, "primary")
        self.cv_start_btn.pack(side=tk.LEFT)
        self.cv_stop_btn = SButton(btn_prog_row, "  停止", self._stop_batch_convert, "danger")
        self.cv_stop_btn.config(state=tk.DISABLED)
        self.cv_stop_btn.pack(side=tk.LEFT, padx=10)
        self.cv_prog = ttk.Progressbar(btn_prog_row, orient=tk.HORIZONTAL, mode="determinate", length=300)
        self.cv_prog.pack(side=tk.LEFT, padx=(10, 0))
        self.cv_prog_lbl = tk.Label(btn_prog_row, text="0 / 0", font=("微软雅黑", 10), bg=C_CARD, fg=C_TEXT)
        self.cv_prog_lbl.pack(side=tk.LEFT, padx=10)

        log_frame = tk.Frame(batch_card, bg=C_CARD)
        log_frame.pack(fill=tk.BOTH, expand=True, padx=14, pady=(0, 6))
        self.cv_log = scrolledtext.ScrolledText(log_frame, font=("Consolas", 10), wrap=tk.WORD, height=20,
                                                 bg="#f8fafc", fg=C_TEXT, relief=tk.SOLID, bd=1, padx=10, pady=10)
        self.cv_log.pack(fill=tk.BOTH, expand=True)

    def _select_cv_batch_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xls")])
        if path:
            self.cv_batch_path.set(path)
            self._load_cv_columns()

    def _load_cv_columns(self):
        path = self.cv_batch_path.get()
        if not path or not os.path.exists(path):
            messagebox.showwarning("提示", "请选择有效的 Excel 文件"); return
        try:
            openpyxl = _get_openpyxl()[0]
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if "Sheet1" not in wb.sheetnames:
                messagebox.showwarning("提示", "Excel 中未找到 Sheet1"); wb.close(); return
            ws = wb["Sheet1"]
            headers = [str(cell.value).strip() if cell.value else f"列{cell.column}" for cell in ws[1]]
            wb.close()
            self.cv_lng_col["values"] = headers
            self.cv_lat_col["values"] = headers
            self.cv_out_lng_col["values"] = headers
            self.cv_out_lat_col["values"] = headers
            lng_match = lat_match = out_lng_match = out_lat_match = None
            for h in headers:
                hl = h.lower()
                if lng_match is None and any(k in hl for k in ["经度", "lng", "lon", "x"]):
                    lng_match = h
                if lat_match is None and any(k in hl for k in ["纬度", "lat", "y"]):
                    lat_match = h
                if out_lng_match is None and any(k in hl for k in ["wgs84经度", "wgs-84经度", "wgs_84经度", "wgs 84经度", "84经度"]):
                    out_lng_match = h
                if out_lat_match is None and any(k in hl for k in ["wgs84纬度", "wgs-84纬度", "wgs_84纬度", "wgs 84纬度", "84纬度"]):
                    out_lat_match = h
            if lng_match: self.cv_lng_col.set(lng_match)
            if lat_match: self.cv_lat_col.set(lat_match)
            if out_lng_match: self.cv_out_lng_col.set(out_lng_match)
            if out_lat_match: self.cv_out_lat_col.set(out_lat_match)
            self.cv_col_frame.pack(fill=tk.X, pady=(0, 10))
            self.cv_log.delete("1.0", tk.END)
            self.cv_log.insert(tk.END, f"已读取文件，共 {len(headers)} 列\n")
            self.cv_log.insert(tk.END, f"表头: {', '.join(headers[:10])}{'...' if len(headers) > 10 else ''}\n")
        except Exception as e:
            messagebox.showerror("错误", f"读取列失败：{e}")

    def _do_single_convert(self):
        self.cv_error.config(text="")
        self.cv_output.config(state="normal"); self.cv_output.delete(0, tk.END); self.cv_output.config(state="readonly")
        try:
            parts = self.cv_input.get().strip().split(",")
            if len(parts) != 2: raise ValueError("格式错误，请输入「经度,纬度」")
            lng, lat = float(parts[0].strip()), float(parts[1].strip())
            trans = self.cv_trans_single.get()
            if "→" not in trans: raise ValueError("请选择有效的转换方式")
            src_name, dst_name = [s.strip() for s in trans.split("→")]
            src_sys = src_name.lower().replace("-", "").replace("_", "")
            dst_sys = dst_name.lower().replace("-", "").replace("_", "")
            src_sys = {"gcj02": "gcj02", "wgs84": "wgs84", "bd09": "bd09"}.get(src_sys, "gcj02")
            dst_sys = {"gcj02": "gcj02", "wgs84": "wgs84", "bd09": "bd09"}.get(dst_sys, "wgs84")
            if src_sys == dst_sys:
                self.cv_error.config(text="源坐标系与目标坐标系相同，无需转换"); return
            out_lng, out_lat = convert_coords(lng, lat, src_sys, dst_sys)
            result_text = f"{round(out_lng, 6)},{round(out_lat, 6)}"
            self.cv_output.config(state="normal"); self.cv_output.delete(0, tk.END)
            self.cv_output.insert(0, result_text); self.cv_output.config(state="readonly")
        except Exception as e:
            self.cv_error.config(text=f"转换失败：{e}")

    def _copy_convert_result(self):
        val = self.cv_output.get()
        if val:
            self.root.clipboard_clear(); self.root.clipboard_append(val)
            self.status_var.set("转换结果已复制")
            self.root.after(2000, lambda: self.status_var.set("就绪"))

    def _on_trans_mode_changed(self, event=None):
        mode = self.cv_trans_mode.get()
        if "→" in mode:
            dst = mode.split("→")[1].strip()
            lng_vals = list(self.cv_out_lng_col["values"] or [])
            lat_vals = list(self.cv_out_lat_col["values"] or [])
            lng_target = f"{dst}经度"
            lat_target = f"{dst}纬度"
            if lng_target in lng_vals: self.cv_out_lng_col.set(lng_target)
            elif lng_vals: self.cv_out_lng_col.set(lng_vals[-1])
            if lat_target in lat_vals: self.cv_out_lat_col.set(lat_target)
            elif lat_vals: self.cv_out_lat_col.set(lat_vals[-1])

    def _stop_batch_convert(self):
        self.stop_flag = True
        self.status_var.set("正在停止...")

    def _start_batch_convert(self):
        path = self.cv_batch_path.get()
        if not path or not os.path.exists(path):
            messagebox.showwarning("提示", "请选择有效的 Excel 文件"); return
        lng_col = self.cv_lng_col.get(); lat_col = self.cv_lat_col.get()
        if not lng_col or not lat_col:
            messagebox.showwarning("提示", "请选择经度列和纬度列"); return
        if lng_col == lat_col:
            messagebox.showwarning("提示", "经度列和纬度列不能相同"); return
        out_lng_name = self.cv_out_lng_col.get().strip()
        out_lat_name = self.cv_out_lat_col.get().strip()
        if not out_lng_name or not out_lat_name:
            messagebox.showwarning("提示", "请选择输出经度列和输出纬度列"); return
        trans_mode = self.cv_trans_mode.get()
        if "→" not in trans_mode:
            messagebox.showwarning("提示", "请选择有效的转换方式"); return
        src_sys_name, dst_sys_name = [s.strip() for s in trans_mode.split("→")]
        src_sys = src_sys_name.lower().replace("-", "").replace("_", "")
        dst_sys = dst_sys_name.lower().replace("-", "").replace("_", "")
        src_sys = {"gcj02": "gcj02", "wgs84": "wgs84", "bd09": "bd09"}.get(src_sys, "gcj02")
        dst_sys = {"gcj02": "gcj02", "wgs84": "wgs84", "bd09": "bd09"}.get(dst_sys, "wgs84")
        if src_sys == dst_sys:
            messagebox.showwarning("提示", "源坐标系与目标坐标系相同，无需转换"); return
        self.stop_flag = False
        self.cv_start_btn.config(state=tk.DISABLED)
        self.cv_stop_btn.config(state=tk.NORMAL)
        self.cv_log.delete("1.0", tk.END)

        def task():
            try:
                openpyxl, PatternFill, Font, Alignment, get_column_letter = _get_openpyxl()
                wb = openpyxl.load_workbook(path)
                if "Sheet1" not in wb.sheetnames:
                    self.root.after(0, lambda: messagebox.showwarning("提示", "Excel 中未找到 Sheet1")); return
                ws = wb["Sheet1"]
                headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
                try:
                    lng_idx = headers.index(lng_col) + 1
                    lat_idx = headers.index(lat_col) + 1
                except ValueError:
                    self.root.after(0, lambda: messagebox.showwarning("提示", "未找到指定的列名")); return
                if out_lng_name in headers:
                    out_lng_idx = headers.index(out_lng_name) + 1
                else:
                    out_lng_idx = len(headers) + 1
                    ws.cell(row=1, column=out_lng_idx, value=out_lng_name)
                    headers.append(out_lng_name)
                if out_lat_name in headers:
                    out_lat_idx = headers.index(out_lat_name) + 1
                else:
                    out_lat_idx = len(headers) + 1
                    ws.cell(row=1, column=out_lat_idx, value=out_lat_name)
                    headers.append(out_lat_name)

                last_row = ws.max_row
                self.root.after(0, lambda: self.cv_prog.config(maximum=last_row - 1, value=0))
                success = fail = 0
                for r in range(2, last_row + 1):
                    if self.stop_flag:
                        self.root.after(0, lambda: self.cv_log.insert(tk.END, "\n用户取消处理\n")); break
                    pv = r - 1
                    self.root.after(0, lambda v=pv, t=last_row - 1: self.cv_prog_lbl.config(text=f"{v} / {t}"))
                    self.root.after(0, lambda v=pv: self.cv_prog.config(value=v))
                    try:
                        lng_val = ws.cell(row=r, column=lng_idx).value
                        lat_val = ws.cell(row=r, column=lat_idx).value
                        if lng_val is None or lat_val is None: continue
                        out_lng, out_lat = convert_coords(float(lng_val), float(lat_val), src_sys, dst_sys)
                        ws.cell(row=r, column=out_lng_idx, value=round(out_lng, 6))
                        ws.cell(row=r, column=out_lat_idx, value=round(out_lat, 6))
                        success += 1
                    except Exception as e:
                        fail += 1
                        ws.cell(row=r, column=out_lng_idx, value=f"错误:{str(e)[:20]}")
                        ws.cell(row=r, column=out_lat_idx, value=f"错误:{str(e)[:20]}")
                wb.save(path)
                self.root.after(0, lambda: self.cv_log.insert(tk.END,
                    f"\n{'='*45}\n转换完成统计\n{'='*45}\n"
                    f"总行数：{last_row - 1}\n成功：{success}\n失败：{fail}\n"
                    f"输出列：{out_lng_name}, {out_lat_name}\n"
                    f"文件已保存：{path}\n{'='*45}\n"))
            except Exception as e:
                self.root.after(0, lambda: self.cv_log.insert(tk.END, f"[处理异常] {str(e)}\n"))
            finally:
                self.root.after(0, lambda: (self.cv_start_btn.config(state=tk.NORMAL),
                                             self.cv_stop_btn.config(state=tk.DISABLED),
                                             self.status_var.set("就绪")))
        threading.Thread(target=task, daemon=True).start()

    # ═══════════════════════════════════════════════════════
    # API 设置
    # ═══════════════════════════════════════════════════════
    def _build_settings(self):
        frame = self._make_page(self.content)
        self.pages["settings"] = frame
        tk.Label(frame, text="API 密钥配置", font=("微软雅黑", 20, "bold"), bg=C_BG, fg=C_TEXT).pack(anchor="w", pady=(0, 20))
        card = Card(frame, title="密钥设置", accent=C_ACCENT)
        card.pack(fill=tk.X, pady=10)
        form = tk.Frame(card, bg=C_CARD)
        form.pack(fill=tk.X, padx=14, pady=14)

        r1 = tk.Frame(form, bg=C_CARD)
        r1.pack(fill=tk.X, pady=8)
        tk.Label(r1, text="* 高德 API Key：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT, width=15, anchor="e").pack(side=tk.LEFT)
        self.cfg_gaode = tk.Entry(r1, font=("微软雅黑", 11), width=55, show="*", relief=tk.SOLID, bd=1)
        self.cfg_gaode.pack(side=tk.LEFT, padx=10, fill=tk.X, expand=True, ipady=2)
        self.cfg_gaode.insert(0, self.config.get("gaode_key", ""))
        tk.Button(r1, text="显示/隐藏", font=("微软雅黑", 10), bg=C_BORDER, fg=C_TEXT, bd=0, padx=10, pady=4,
                 command=lambda: self.cfg_gaode.config(show="" if self.cfg_gaode.cget("show") == "*" else "*")).pack(side=tk.LEFT)

        r2 = tk.Frame(form, bg=C_CARD)
        r2.pack(fill=tk.X, pady=8)
        tk.Label(r2, text="  百度 AK：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT, width=15, anchor="e").pack(side=tk.LEFT)
        self.cfg_baidu = tk.Entry(r2, font=("微软雅黑", 11), width=55, show="*", relief=tk.SOLID, bd=1)
        self.cfg_baidu.pack(side=tk.LEFT, padx=10, fill=tk.X, expand=True, ipady=2)
        self.cfg_baidu.insert(0, self.config.get("baidu_ak", ""))
        tk.Button(r2, text="显示/隐藏", font=("微软雅黑", 10), bg=C_BORDER, fg=C_TEXT, bd=0, padx=10, pady=4,
                 command=lambda: self.cfg_baidu.config(show="" if self.cfg_baidu.cget("show") == "*" else "*")).pack(side=tk.LEFT)

        opt = tk.Frame(form, bg=C_CARD)
        opt.pack(fill=tk.X, pady=8)
        self.cfg_enable_baidu = tk.BooleanVar(value=self.config.get("enable_baidu_fallback", True))
        tk.Checkbutton(opt, text="启用百度地图 Fallback（高德精度不足时自动切换）", variable=self.cfg_enable_baidu,
                      font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT, selectcolor=C_CARD).pack(anchor="w")

        r3 = tk.Frame(form, bg=C_CARD)
        r3.pack(fill=tk.X, pady=8)
        tk.Label(r3, text="百度触发等级：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT, width=15, anchor="e").pack(side=tk.LEFT)
        self.cfg_trigger = ttk.Combobox(r3, values=["门址", "街道", "乡镇", "区县", "城市"], font=("微软雅黑", 10), width=15, state="readonly")
        self.cfg_trigger.set(self.config.get("baidu_trigger_level", "乡镇"))
        self.cfg_trigger.pack(side=tk.LEFT, padx=10)
        tk.Label(r3, text="请求间隔（秒）：", font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT).pack(side=tk.LEFT, padx=(30, 0))
        self.cfg_interval = tk.Spinbox(r3, from_=0.1, to=5.0, increment=0.1, width=8, font=("微软雅黑", 10), relief=tk.SOLID, bd=1)
        self.cfg_interval.delete(0, tk.END)
        self.cfg_interval.insert(0, str(self.config.get("base_interval", 0.2)))
        self.cfg_interval.pack(side=tk.LEFT, padx=10)

        bf = tk.Frame(card, bg=C_CARD)
        bf.pack(fill=tk.X, padx=14, pady=(0, 14))
        SButton(bf, "  保存配置", self.save_settings, "primary").pack(side=tk.LEFT)
        SButton(bf, "  测试高德", lambda: self.test_api("gaode"), "success").pack(side=tk.LEFT, padx=10)
        SButton(bf, "  测试百度", lambda: self.test_api("baidu"), "success").pack(side=tk.LEFT)

        tip = Card(frame, title="使用说明", accent=C_INFO)
        tip.pack(fill=tk.X, pady=16)
        for t in [
            "• 高德 Key 为必填项，请前往 https://console.amap.com 申请个人开发者 Key",
            "• 百度 AK 为可选项，用于提升解析精度，前往 https://lbsyun.baidu.com 申请",
            "• 配置将自动保存到当前目录的 addr_config.json 文件中，下次打开自动加载",
            "• 高德个人开发者每天约有 5000 次免费地理编码额度，超出需付费",
            "• 建议开启百度 Fallback，当高德返回精度低于「乡镇」时自动尝试百度"
        ]:
            tk.Label(tip, text=t, font=("微软雅黑", 11), bg=C_CARD, fg=C_TEXT_M).pack(anchor="w", padx=15, pady=5)

    def save_settings(self):
        self.config["gaode_key"] = self.cfg_gaode.get().strip()
        self.config["baidu_ak"] = self.cfg_baidu.get().strip()
        self.config["enable_baidu_fallback"] = self.cfg_enable_baidu.get()
        self.config["baidu_trigger_level"] = self.cfg_trigger.get()
        try: self.config["base_interval"] = float(self.cfg_interval.get())
        except: self.config["base_interval"] = 0.2
        save_config(self.config)
        messagebox.showinfo("提示", "配置已保存！")

    def test_api(self, api_type):
        key = self.cfg_gaode.get().strip() if api_type == "gaode" else self.cfg_baidu.get().strip()
        if not key: messagebox.showwarning("提示", "请先输入 API Key"); return
        self.status_var.set(f"正在测试 {api_type} ...")
        def task():
            if api_type == "gaode":
                url = f"{GAODE_GEO}?address={url_enc('北京市天安门')}&output=json&key={key}"
                ok, resp = http_get(url, 2)
                if ok and json_val(resp, "status") == "1":
                    self.root.after(0, lambda: messagebox.showinfo("测试成功", "高德 API 连接正常！"))
                else:
                    code = json_val(resp, "infocode")
                    msg = get_gaode_error_msg(code, json_val(resp, "info") if ok else resp)
                    self.root.after(0, lambda: messagebox.showerror("测试失败", f"错误：{msg}"))
            else:
                url = f"{BAIDU_GEO}?address={url_enc('北京市天安门')}&output=json&ak={key}"
                ok, resp = http_get(url, 2)
                if ok and json_val(resp, "status") == "0":
                    self.root.after(0, lambda: messagebox.showinfo("测试成功", "百度 API 连接正常！"))
                else:
                    code = json_val(resp, "status")
                    msg = get_baidu_error_msg(code, resp if ok else resp)
                    self.root.after(0, lambda: messagebox.showerror("测试失败", f"错误：{msg}"))
            self.root.after(0, lambda: self.status_var.set("就绪"))
        threading.Thread(target=task, daemon=True).start()


# ═══════════════════════════════════════════════════════
# 入口
# ═══════════════════════════════════════════════════════
def main():
    root = tk.Tk()
    app = App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
